"""Reproducible end-to-end demo; the user source template is never overwritten."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from pptx import Presentation

from genppt import GeneratorConfig, generate_presentation


def prepare_template(source: Path, target: Path) -> None:
    p = Presentation(source)
    if len(p.slides) != 3:
        raise ValueError("this demo expects the three-page CPS/CNS/overall source")
    first = next(s for s in p.slides[0].shapes if s.has_table)
    first.name = "<foreach item=cpsTables>CPS table</foreach>"
    group = p.slides[1]
    caption = next(s for s in group.shapes if s.has_text_frame and "<foreach" in s.name)
    caption.name = "<foreach item=tables block=reasonTable>Classification</foreach>"
    paragraph = caption.text_frame.paragraphs[0]
    # Preserve the reference font properties when changing test content.
    if paragraph.runs:
        paragraph.runs[0].text = "分类：{{title}}<if status> — {{status}}</if>"
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = "分类：{{title}}<if status> — {{status}}</if>"
    native = next(s for s in group.shapes if s.has_table)
    native.name = "<foreach item=tables block=reasonTable>{{table1}}</foreach>"
    pictures = [s for s in group.shapes if s.shape_type == 13]
    if len(pictures) != 2:
        raise ValueError("CNS demo source must contain the two reference pictures")
    pictures[0].name = "<foreach item=tables block=reasonTable cols=1>{{images}}</foreach>（x，y）"
    pictures[1].name = "<foreach item=tables block=reasonTable>{{image2}}</foreach>（x，y）"
    last = next(s for s in p.slides[2].shapes if s.has_table)
    last.name = "<foreach item=ovrTables>Overall table</foreach>"
    target.parent.mkdir(parents=True, exist_ok=True)
    p.save(target)


def demo_payload() -> dict:
    images = Path(__file__).resolve().parents[1] / "tests/fixtures/images"
    return {
        "cpsTables": [
            {
                "columns": [{"loc": "Region"}, {"count": "Count"}],
                "rows": [{"loc": "Central", "count": 12}, {"loc": "South", "count": 8}],
            },
            {"rows": [{"reason": "Timeout", "count": 3}]},
        ],
        "tables": [
            {
                "title": "Long equipment report",
                "status": "Normal",
                "table1": {
                    "columns": [{"id": "ID"}, {"value": "Value"}],
                    "rows": [{"id": i, "value": f"Row {i}"} for i in range(1, 26)],
                },
                "images": [str(images / "wide.png"), str(images / "square.png")],
                "image2": str(images / "tall.png"),
            },
            {
                "title": "Second group",
                "table1": [
                    {"rows": [{"price": 3000, "amount": 6000}]},
                    {"columns": [{"note": "Note"}], "rows": [{"note": "Second inner table"}]},
                ],
                "images": "",
                "image2": str(images / "square.png"),
            },
            {"title": "Text only", "status": "No photos"},
        ],
        "ovrTables": [{"rows": [{"category": "Overall", "count": 42}]}],
    }


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=root / "tests/fixtures/test-multipage.pptx")
    parser.add_argument("--output-dir", type=Path, default=root / "output/multipage-demo")
    parser.add_argument("--prepare", action="store_true")
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    template = args.source
    if args.prepare:
        template = args.output_dir / "template.pptx"
        prepare_template(args.source, template)
    payload = demo_payload()
    (args.output_dir / "data.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    result = generate_presentation(
        template,
        payload,
        args.output_dir / "generated.pptx",
        config=GeneratorConfig(allowed_image_roots=(str(root / "tests/fixtures/images"),)),
    )
    print(result)


if __name__ == "__main__":
    main()
