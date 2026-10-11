from __future__ import annotations

from collections.abc import Mapping
from copy import deepcopy
from dataclasses import dataclass
from typing import Any

from pptx.text.text import _Paragraph

from ..core.config import GeneratorConfig
from ..core.errors import TemplateError
from .data import normalize_table
from .directives import optional_path, select_value
from .image_renderer import ImageCache, _picture_size, _replace_picture, _validate_image
from .models import ShapePrototype
from .planner import PageBlock, PagePlan
from .pptx_utils import append_element, clone_blank_slide, materialize_text_shape, remove_shape
from .table_renderer import add_table_fragment, prepare_table, split_column_indices
from .text_renderer import TextBudget, paragraph_height, render_text_frame, text_height


@dataclass
class Resources:
    config: GeneratorConfig
    text: TextBudget
    images: int = 0
    tables: int = 0
    cells: int = 0
    output_cells: int = 0

    def add_images(self, count: int) -> None:
        self.images += count
        if self.images > self.config.max_images:
            raise TemplateError("image count exceeds configured limit")


@dataclass
class Unit:
    kind: str
    prototype: ShapePrototype
    payload: Any
    left: int
    width: int
    height: int
    offset: int = 0
    break_before: bool = False


@dataclass
class Band:
    lanes: list[list[Unit]]
    gap: int

    @property
    def height(self) -> int:
        return max(sum(unit.height for unit in lane) + lane[0].offset for lane in self.lanes)


@dataclass
class Placement:
    unit: Unit
    page: int
    top: int
    height: int
    start: int = 0
    count: int = 0


def _array(value: Any, *, scalar: bool = False) -> list:
    if value is None or value == "" or value == {}:
        return []
    if isinstance(value, list):
        return value
    if scalar:
        return [value]
    raise TemplateError("foreach item must resolve to an array")


def _image_units(
    prototype: ShapePrototype,
    raw: Any,
    right: int,
    usable_height: int,
    resources: Resources,
    cache: ImageCache,
) -> list[Unit]:
    values = _array(raw, scalar=True)
    resources.add_images(len(values))
    infos = [
        info
        for value in values
        if (info := _validate_image(value, resources.config, cache)) is not None
    ]
    if not infos:
        return []
    shape, directive = prototype.shape, prototype.directive
    sizes = [_picture_size(shape, info, directive.image_mode) for info in infos]
    width, height = max(w for w, _ in sizes), max(h for _, h in sizes)
    config = resources.config
    gap = config.block_gap
    available_width = right - shape.left
    cols, rows = directive.cols or 1, directive.rows or 1
    scale = min(
        1.0,
        (available_width - gap * (cols - 1)) / (width * cols),
        (usable_height - gap * (rows - 1)) / (height * rows),
    )
    if scale <= 0:
        raise TemplateError("image grid has no usable cell size")
    cell_width, cell_height = max(1, round(width * scale)), max(1, round(height * scale))
    if directive.cols is None:
        cols = max(1, (available_width + gap) // (cell_width + gap))
    if directive.rows is None:
        rows = max(1, (usable_height + gap) // (cell_height + gap))
    units = []
    for row_index, start in enumerate(range(0, len(infos), cols)):
        pictures = []
        for col, index in enumerate(range(start, min(start + cols, len(infos)))):
            w, h = sizes[index]
            pictures.append(
                (
                    infos[index],
                    col * (cell_width + gap),
                    max(1, round(w * scale)),
                    max(1, round(h * scale)),
                )
            )
        units.append(
            Unit(
                "images",
                prototype,
                pictures,
                shape.left,
                available_width,
                cell_height,
                break_before=(row_index > 0 and row_index % rows == 0),
            )
        )
        if start + cols < len(infos):
            units.append(Unit("gap", prototype, None, shape.left, 0, gap))
    return units


def _table_units(prototype: ShapePrototype, raw: Any, resources: Resources) -> list[Unit]:
    units = []
    source = prototype.shape.table
    if any(c.is_merge_origin or c.is_spanned for c in source.iter_cells()):
        raise TemplateError("merged cells are not supported in a dynamic table prototype")
    for item in _array(raw, scalar=True):
        if not isinstance(item, Mapping):
            raise TemplateError("table data must be an object or object array")
        resources.tables += 1
        if resources.tables > resources.config.max_tables:
            raise TemplateError("table count exceeds configured limit")
        raw_rows = item.get("rows")
        raw_columns = item.get("columns")
        if isinstance(raw_rows, list) and raw_rows:
            column_count = (
                len(raw_columns)
                if isinstance(raw_columns, list)
                else len(raw_rows[0])
                if isinstance(raw_rows[0], Mapping)
                else 0
            )
            if (
                resources.cells + column_count * (len(raw_rows) + 1)
                > resources.config.max_total_cells
            ):
                raise TemplateError("total cell count exceeds configured limit")
        normalized = normalize_table(item, resources.config)
        if not normalized.rows:
            continue
        resources.cells += len(normalized.labels) * (len(normalized.rows) + 1)
        if resources.cells > resources.config.max_total_cells:
            raise TemplateError("total cell count exceeds configured limit")
        groups = split_column_indices(
            len(normalized.labels),
            prototype.shape.width,
            normalized.options.min_column_width_inches or resources.config.min_column_width_inches,
            normalized.options.repeat_leading_columns,
        )
        for columns in groups:
            prepared = prepare_table(source, normalized, columns)
            if units:
                units.append(
                    Unit(
                        "gap", prototype, None, prototype.shape.left, 0, resources.config.block_gap
                    )
                )
            units.append(
                Unit(
                    "table",
                    prototype,
                    prepared,
                    prototype.shape.left,
                    prototype.shape.width,
                    sum(prepared.row_heights),
                )
            )
    return units


def _text_units(prototype: ShapePrototype, context: Mapping, resources: Resources) -> list[Unit]:
    from pptx.shapes.autoshape import Shape

    # The detached clone preserves styles without changing the template prototype.
    shape = Shape(materialize_text_shape(prototype.shape), prototype.shape._parent)
    if not shape.has_text_frame:
        return [Unit("shape", prototype, shape._element, shape.left, shape.width, shape.height)]
    render_text_frame(shape.text_frame, context, resources.text)
    if not shape.text.strip():
        return []
    return [Unit("text", prototype, shape._element, shape.left, shape.width, text_height(shape))]


def _prototype_bands(block: PageBlock) -> list[list[ShapePrototype]]:
    bands: list[list[ShapePrototype]] = []
    for prototype in sorted(block.prototypes, key=lambda p: (p.shape.top, p.shape.left)):
        if bands:
            previous = bands[-1]
            same_row = prototype.shape.top < min(p.shape.top + p.shape.height for p in previous)
            disjoint = all(
                prototype.shape.left >= p.shape.left + p.shape.width
                or p.shape.left >= prototype.shape.left + prototype.shape.width
                for p in previous
            )
            if same_row and disjoint:
                previous.append(prototype)
                continue
        bands.append([prototype])
    return bands


def _prepare_group(
    block: PageBlock,
    context: Any,
    right: int,
    full_height: int,
    resources: Resources,
    cache: ImageCache,
) -> list[Band]:
    bands = []
    previous_bottom = block.top
    for prototypes in _prototype_bands(block):
        top = min(p.shape.top for p in prototypes)
        gap = max(0, top - previous_bottom)
        lanes = []
        for prototype in prototypes:
            shape = prototype.shape
            if shape.has_table:
                table_data = select_value(prototype.directive.body, context)
                if (
                    isinstance(context, Mapping)
                    and "rows" in context
                    and (
                        table_data is not None
                        or prototype.directive.body in {"{{tableName}}", "{{title}}"}
                    )
                    and not isinstance(table_data, (Mapping, list))
                ):
                    table_data = context
                units = _table_units(prototype, table_data, resources)
            elif hasattr(shape, "_pic"):
                # A neighboring lane bounds automatic image-array grid expansion.
                lane_right = min(
                    [right]
                    + [
                        p.shape.left - resources.config.block_gap
                        for p in prototypes
                        if p.shape.left > shape.left
                    ]
                )
                units = _image_units(
                    prototype,
                    select_value(prototype.directive.body, context),
                    lane_right,
                    full_height - (shape.top - top),
                    resources,
                    cache,
                )
            else:
                if not isinstance(context, Mapping):
                    raise TemplateError("text/group foreach requires an object array")
                units = _text_units(prototype, context, resources)
            if units:
                units[0].offset = shape.top - top
                lanes.append(units)
        if lanes:
            bands.append(Band(lanes, gap if bands else 0))
        previous_bottom = max(p.shape.top + p.shape.height for p in prototypes)
    return bands


def _layout_lane(
    lane: list[Unit],
    page: int,
    cursor: int,
    first_bottom: int,
    fresh_top: int,
    fresh_bottom: int,
    minimum_fragment_rows: int,
) -> tuple[list[Placement], int, int]:
    placements = []

    def next_page() -> None:
        nonlocal page, cursor
        page += 1
        cursor = fresh_top

    for unit in lane:
        bottom = first_bottom if page == 0 else fresh_bottom
        if unit.break_before and cursor > fresh_top:
            next_page()
            bottom = fresh_bottom
        if unit.kind == "gap":
            cursor = min(bottom, cursor + unit.height)
            continue
        if unit.height <= bottom - cursor:
            count = len(unit.payload.rows) if unit.kind == "table" else 0
            placements.append(Placement(unit, page, cursor, unit.height, 0, count))
            cursor += unit.height
            continue
        # Unsplittable units first try a full fresh page.
        if unit.height <= fresh_bottom - fresh_top:
            next_page()
            count = len(unit.payload.rows) if unit.kind == "table" else 0
            placements.append(Placement(unit, page, cursor, unit.height, 0, count))
            cursor += unit.height
            continue
        if unit.kind == "table":
            heights = unit.payload.row_heights
            start = 0
            while start < len(unit.payload.rows):
                bottom = first_bottom if page == 0 else fresh_bottom
                height, count = heights[0], 0
                while start + count < len(unit.payload.rows):
                    row_height = heights[1 + start + count]
                    if height + row_height > bottom - cursor:
                        break
                    height += row_height
                    count += 1
                if count == 0:
                    if cursor == fresh_top and page > 0:
                        raise TemplateError(
                            "table header and data row cannot fit in continuation region"
                        )
                    next_page()
                    continue
                if (
                    count < min(minimum_fragment_rows, len(unit.payload.rows) - start)
                    and cursor > fresh_top
                ):
                    # Avoid a tiny initial fragment when a full fresh page can
                    # accommodate more rows. On a fresh page this is advisory:
                    # a valid oversized row must still make forward progress.
                    next_page()
                    continue
                placements.append(Placement(unit, page, cursor, height, start, count))
                cursor += height
                start += count
                if start < len(unit.payload.rows):
                    next_page()
        elif unit.kind == "text":
            from pptx.shapes.autoshape import Shape

            # Fragment measurement must use the same materialized inheritance
            # as the rendered text, not the placeholder's original layout defaults.
            shape = Shape(unit.payload, unit.prototype.shape._parent)
            body = unit.payload.find(
                "{http://schemas.openxmlformats.org/presentationml/2006/main}txBody"
            )
            paragraphs = body.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}p")
            frame = shape.text_frame
            margin = frame.margin_top + frame.margin_bottom
            width = unit.width - frame.margin_left - frame.margin_right
            for index, paragraph in enumerate(paragraphs):
                height = paragraph_height(_Paragraph(paragraph, frame), width) + margin
                bottom = first_bottom if page == 0 else fresh_bottom
                if height > fresh_bottom - fresh_top:
                    raise TemplateError("single text paragraph cannot fit in continuation region")
                if height > bottom - cursor:
                    next_page()
                placements.append(Placement(unit, page, cursor, height, index, 1))
                cursor += height
        else:
            raise TemplateError("shape cannot fit in continuation region")
    return placements, page, cursor


def _layout_group(
    bands: list[Band],
    page: int,
    cursor: int,
    first_bottom: int,
    fresh_top: int,
    fresh_bottom: int,
    minimum_fragment_rows: int,
) -> tuple[list[Placement], int, int]:
    result = []
    for band in bands:
        cursor += band.gap
        bottom = first_bottom if page == 0 else fresh_bottom
        if band.height > bottom - cursor and band.height <= fresh_bottom - fresh_top:
            page += 1
            cursor = fresh_top
        endings = []
        for lane in band.lanes:
            placements, end_page, end_cursor = _layout_lane(
                lane,
                page,
                cursor + lane[0].offset,
                first_bottom,
                fresh_top,
                fresh_bottom,
                minimum_fragment_rows,
            )
            result.extend(placements)
            endings.append((end_page, end_cursor))
        page = max(end_page for end_page, _ in endings)
        cursor = max(end_cursor for end_page, end_cursor in endings if end_page == page)
    return result, page, cursor


def _render_placement(
    placement: Placement,
    slide: Any,
    source_slide: Any,
    name: str,
) -> None:
    unit = placement.unit
    if unit.kind == "table":
        add_table_fragment(
            slide,
            unit.prototype,
            unit.payload,
            unit.left,
            placement.top,
            placement.start,
            placement.count,
            name,
        )
    elif unit.kind == "images":
        for index, (info, offset, width, height) in enumerate(unit.payload):
            shape = append_element(slide, unit.prototype.element, source_slide)
            shape.left, shape.top = unit.left + offset, placement.top
            _replace_picture(slide, shape, info, unit.prototype.directive.image_mode)
            shape.width, shape.height = width, height
            shape.name = f"{name}_image{index + 1}"
    else:
        element = deepcopy(unit.payload)
        if unit.kind == "text" and placement.count:
            body = element.find(
                "{http://schemas.openxmlformats.org/presentationml/2006/main}txBody"
            )
            paragraphs = body.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}p")
            for index, paragraph in enumerate(paragraphs):
                if index != placement.start:
                    body.remove(paragraph)
        shape = append_element(slide, element, source_slide)
        shape.left, shape.top, shape.height = unit.left, placement.top, placement.height
        shape.name = name


def expand_pages(
    presentation: Any,
    plans: list[PagePlan],
    data: Mapping,
    resources: Resources,
    cache: ImageCache,
) -> None:
    config = resources.config
    fresh_top = config.continuation_top_margin
    bottom = presentation.slide_height - config.bottom_margin
    if bottom <= fresh_top:
        raise TemplateError("continuation margins leave no usable region")
    for plan in plans:
        source = plan.slide
        for shape in list(source.shapes):
            if shape.shape_id in plan.dynamic_ids:
                remove_shape(shape)
        insert_after = source
        for block in plan.blocks:
            first_bottom = min(
                [bottom]
                + [
                    other.top - config.block_gap
                    for other in plan.blocks
                    if other.top >= block.bottom
                    and other.left < block.right
                    and block.left < other.right
                ]
            )
            right = min(
                [presentation.slide_width - config.bottom_margin]
                + [
                    other.left - config.block_gap
                    for other in plan.blocks
                    if other.left >= block.right
                    and other.top < block.bottom
                    and block.top < other.bottom
                ]
            )
            # Template geometry outside the slide is an error, not silently cropped.
            for p in block.prototypes:
                if p.shape.left < 0 or p.shape.left + p.shape.width > presentation.slide_width:
                    raise TemplateError("dynamic Shape extends beyond slide width")
            raw = data if block.root_text else optional_path(data, block.item_path)
            items = (
                [raw]
                if block.root_text
                else _array(
                    raw,
                    scalar=(
                        len(block.prototypes) == 1
                        and (
                            block.prototypes[0].shape.has_table
                            or hasattr(block.prototypes[0].shape, "_pic")
                        )
                    ),
                )
            )
            resources.text.loop_items += len(items)
            if resources.text.loop_items > config.max_loop_items:
                raise TemplateError("loop item count exceeds configured limit")
            slides = [source]
            page, cursor = 0, block.top
            rendered_items = 0
            # Standalone images form one grid across the outer path/object array.
            standalone_image = len(block.prototypes) == 1 and hasattr(
                block.prototypes[0].shape, "_pic"
            )
            if standalone_image and block.page != "group":
                prototype = block.prototypes[0]
                paths = []
                for item in items:
                    value = select_value(prototype.directive.body, item)
                    paths.extend(_array(value, scalar=True))
                items = [paths]
            for item_index, item in enumerate(items):
                if standalone_image and block.page != "group":
                    units = _image_units(
                        block.prototypes[0],
                        item,
                        right,
                        min(bottom - fresh_top, first_bottom - block.top),
                        resources,
                        cache,
                    )
                    bands = [Band([units], 0)] if units else []
                else:
                    bands = _prepare_group(block, item, right, bottom - fresh_top, resources, cache)
                if not bands:
                    continue
                total_height = sum(b.height + b.gap for b in bands)
                current_bottom = first_bottom if page == 0 else bottom
                if (block.page == "group" and rendered_items) or (
                    total_height > current_bottom - cursor
                    and (rendered_items or total_height <= bottom - fresh_top)
                ):
                    page += 1
                    cursor = fresh_top
                placements, page, cursor = _layout_group(
                    bands,
                    page,
                    cursor,
                    first_bottom,
                    fresh_top,
                    bottom,
                    config.minimum_fragment_rows,
                )
                if page + 1 > config.max_output_slides:
                    raise TemplateError("output slide count exceeds limit")
                while len(slides) <= page:
                    if len(presentation.slides) >= config.max_output_slides:
                        raise TemplateError("output slide count exceeds limit")
                    slide = clone_blank_slide(presentation, source, insert_after)
                    insert_after = slide
                    slides.append(slide)
                for fragment_index, placement in enumerate(placements):
                    if placement.unit.kind == "table":
                        resources.output_cells += len(placement.unit.payload.labels) * (
                            placement.count + 1
                        )
                        if resources.output_cells > config.max_total_cells:
                            raise TemplateError("generated cell count exceeds configured limit")
                    base = placement.unit.prototype.directive.body or placement.unit.kind
                    if placement.unit.kind not in {"text", "shape"}:
                        selected = select_value(base, item)
                        if isinstance(selected, (str, int)):
                            base = str(selected)
                        else:
                            base = placement.unit.kind
                    suffix = (
                        f"__{block.block_name[:40]}_i{item_index + 1}"
                        f"_p{plan.source_index + 1}_s{placement.unit.prototype.shape.shape_id}"
                        f"_f{fragment_index + 1}"
                    )
                    name = base[: max(1, 240 - len(suffix))] + suffix
                    _render_placement(placement, slides[placement.page], source, name)
                cursor += config.block_gap
                rendered_items += 1
