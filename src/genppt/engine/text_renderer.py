from __future__ import annotations

import re
from bisect import bisect_right
from collections.abc import Mapping
from copy import deepcopy
from dataclasses import dataclass
from typing import Any

from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

from ..core.config import GeneratorConfig
from ..core.errors import TemplateError
from .directives import EXPRESSION_RE, optional_path, parse_foreach

PATH = r"[A-Za-z_][\w.]*"
TOKEN_RE = re.compile(r"{{\s*(" + PATH + r")\s*}}|<if\s+(" + PATH + r")>|</if>")
LOOP_RE = re.compile(
    r"^\s*<foreach\s+items\s*=\s*[\"']?(" + PATH + r")[\"']?\s*>(.*?)</foreach>\s*$", re.DOTALL
)


@dataclass
class TextBudget:
    config: GeneratorConfig
    paragraphs: int = 0
    characters: int = 0
    loop_items: int = 0

    def add(self, text: str) -> None:
        self.paragraphs += 1
        self.characters += len(text)
        if self.paragraphs > self.config.max_text_paragraphs:
            raise TemplateError("generated text paragraph count exceeds limit")
        if self.characters > self.config.max_text_characters:
            raise TemplateError("generated text character count exceeds limit")


def _source(paragraph: Any) -> tuple[str, list[tuple[int, int, Any]]]:
    text = ""
    spans = []
    for node in paragraph._p:
        if node.tag.endswith("}r") or node.tag.endswith("}fld"):
            value = "".join(node.xpath("./a:t/text()"))
        elif node.tag.endswith("}br"):
            value = "\v"
        else:
            continue
        spans.append((len(text), len(text) + len(value), node))
        text += value
    return text, spans


def _compile_inline(text: str) -> list[tuple[str, int, int, str | None]]:
    tokens = []
    depth = 0
    for match in TOKEN_RE.finditer(text):
        if match.group(1):
            kind, value = "value", match.group(1)
        elif match.group(2):
            kind, value = "if", match.group(2)
            depth += 1
            if depth > 16:
                raise TemplateError("if nesting exceeds limit 16")
        else:
            kind, value = "end", None
            depth -= 1
            if depth < 0:
                raise TemplateError("unmatched </if> in text")
        tokens.append((kind, match.start(), match.end(), value))
    if depth:
        raise TemplateError("unclosed <if> in text")
    remaining = TOKEN_RE.sub("", text)
    if re.search(r"</?(?:if|foreach)\b", remaining):
        raise TemplateError("invalid text tag or foreach must occupy an entire paragraph")
    return tokens


def validate_text_frame(text_frame: Any) -> None:
    for paragraph in text_frame.paragraphs:
        text, _ = _source(paragraph)
        loop = LOOP_RE.fullmatch(text)
        _compile_inline(loop.group(2) if loop else text)


def _slice(spans: list[tuple[int, int, Any]], start: int, end: int):
    index = bisect_right(spans, start, key=lambda span: span[1])
    while index < len(spans):
        low, high, node = spans[index]
        if low >= end:
            break
        if high > start and low < end:
            original = "\v" if node.tag.endswith("}br") else "".join(node.xpath("./a:t/text()"))
            yield original[max(0, start - low) : min(high, end) - low], node
        index += 1


def _evaluate(
    text: str, spans: list, start: int, end: int, context: Mapping[str, Any]
) -> list[tuple[str, Any]]:
    tokens = _compile_inline(text[start:end])
    pieces = []
    active = [True]
    cursor = start
    for kind, low, high, path in tokens:
        low, high = low + start, high + start
        if active[-1]:
            pieces.extend(_slice(spans, cursor, low))
        if kind == "if":
            active.append(active[-1] and bool(optional_path(context, path)))
        elif kind == "end":
            active.pop()
        elif active[-1]:
            value = optional_path(context, path)
            value = "" if value is None else str(value)
            index = bisect_right(spans, low, key=lambda span: span[1])
            style = spans[index][2] if index < len(spans) else None
            pieces.append((value, style))
        cursor = high
    if active[-1]:
        pieces.extend(_slice(spans, cursor, end))
    return pieces


def _paragraph_from(source: Any, pieces: list[tuple[str, Any]]) -> Any:
    result = deepcopy(source._p)
    for node in list(result):
        if not node.tag.endswith("}pPr") and not node.tag.endswith("}endParaRPr"):
            result.remove(node)
    end = result.find("{http://schemas.openxmlformats.org/drawingml/2006/main}endParaRPr")
    for text, style in pieces:
        if not text:
            continue
        if text == "\v" and style is not None and style.tag.endswith("}br"):
            run = deepcopy(style)
        elif (
            style is not None
            and style.tag.endswith("}fld")
            and text == "".join(style.xpath("./a:t/text()"))
        ):
            # An untouched date/page field beside a placeholder stays a field.
            run = deepcopy(style)
        else:
            run = OxmlElement("a:r")
            if style is not None:
                props = style.find("{http://schemas.openxmlformats.org/drawingml/2006/main}rPr")
                if props is not None:
                    run.append(deepcopy(props))
            node = OxmlElement("a:t")
            node.text = text
            run.append(node)
        if end is None:
            result.append(run)
        else:
            end.addprevious(run)
    return result


def render_text_frame(
    text_frame: Any, context: Mapping[str, Any], budget: TextBudget | None = None
) -> int:
    output = []
    count = 0
    for paragraph in text_frame.paragraphs:
        text, spans = _source(paragraph)
        loop = LOOP_RE.fullmatch(text)
        if not loop and not TOKEN_RE.search(text):
            # Preserve dynamic fields (page numbers, dates), hyperlinks and
            # unsupported run nodes in untouched template paragraphs verbatim.
            if budget:
                budget.add(text)
            output.append(deepcopy(paragraph._p))
            continue
        contexts = [context]
        start, end = 0, len(text)
        if loop:
            start, end = loop.start(2), loop.end(2)
            values = optional_path(context, loop.group(1))
            if values is None:
                values = []
            if not isinstance(values, list) or any(not isinstance(v, Mapping) for v in values):
                raise TemplateError("paragraph foreach requires an object array")
            if budget:
                budget.loop_items += len(values)
                if budget.loop_items > budget.config.max_loop_items:
                    raise TemplateError("text loop item count exceeds limit")
            contexts = values
        for scope in contexts:
            pieces = _evaluate(text, spans, start, end, scope)
            rendered = "".join(value for value, _ in pieces)
            if (loop or "<if" in text) and not rendered.strip():
                continue
            if budget:
                budget.add(rendered)
            output.append(_paragraph_from(paragraph, pieces))
        count += len(EXPRESSION_RE.findall(text))
    for paragraph in list(text_frame._txBody.p_lst):
        text_frame._txBody.remove(paragraph)
    if not output:
        output = [OxmlElement("a:p")]
    for paragraph in output:
        text_frame._txBody.append(paragraph)
    return count


def replace_paragraph_expressions(paragraph: Any, context: Mapping[str, Any]) -> int:
    text, spans = _source(paragraph)
    pieces = _evaluate(text, spans, 0, len(text), context)
    result = _paragraph_from(paragraph, pieces)
    paragraph._p.getparent().replace(paragraph._p, result)
    return len(EXPRESSION_RE.findall(text))


def replace_text_frame_expressions(text_frame: Any, context: Mapping[str, Any]) -> int:
    return render_text_frame(text_frame, context)


def paragraph_height(paragraph: Any, width: int) -> int:
    from .table_renderer import _line_count

    size = max(
        [run.font.size.pt for run in paragraph.runs if run.font.size]
        + [paragraph.font.size.pt if paragraph.font.size else 12.0]
    )
    lines = _line_count(paragraph.text.replace("\v", "\n"), max(1, width / 12700), size)
    spacing = paragraph.line_spacing
    line_height = spacing.pt if hasattr(spacing, "pt") else size * (spacing or 1.22)
    before = paragraph.space_before or 0
    after = paragraph.space_after or 0
    return int(Pt(line_height * lines) + before + after)


def text_height(shape: Any) -> int:
    frame = shape.text_frame
    width = shape.width - frame.margin_left - frame.margin_right
    return max(
        shape.height,
        frame.margin_top
        + frame.margin_bottom
        + sum(paragraph_height(p, width) for p in frame.paragraphs),
    )


def _process_shape(shape: Any, context: Mapping[str, Any], budget: TextBudget) -> int:
    if parse_foreach(shape.name, shape.shape_id) is not None:
        return 0
    count = 0
    if getattr(shape, "has_text_frame", False):
        count += render_text_frame(shape.text_frame, context, budget)
    if getattr(shape, "has_table", False):
        for cell in shape.table.iter_cells():
            count += render_text_frame(cell.text_frame, context, budget)
    if hasattr(shape, "shapes"):
        for child in shape.shapes:
            count += _process_shape(child, context, budget)
    return count


def process_text_expressions(
    presentation: Any, data: Mapping[str, Any], budget: TextBudget | None = None
) -> int:
    budget = budget or TextBudget(GeneratorConfig())
    return sum(
        _process_shape(shape, data, budget)
        for slide in presentation.slides
        for shape in slide.shapes
    )
