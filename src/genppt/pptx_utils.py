from __future__ import annotations

from copy import deepcopy
from typing import Any

from pptx.oxml.xmlchemy import OxmlElement


def _merge_properties(base: Any, overlay: Any) -> Any:
    result = deepcopy(base) if base is not None else OxmlElement("a:pPr")
    if overlay is None:
        return result
    result.attrib.update(overlay.attrib)
    for child in overlay:
        local = child.tag.rsplit("}", 1)[-1]
        choices = (
            {"buNone", "buChar", "buAutoNum", "buBlip"},
            {"buClrTx", "buClr"},
            {"buSzTx", "buSzPct", "buSzPts"},
            {"buFontTx", "buFont"},
            {"solidFill", "noFill", "gradFill", "pattFill", "blipFill", "grpFill"},
            {"noAutofit", "normAutofit", "spAutoFit"},
            {"effectLst", "effectDag"},
            {"uLnTx", "uLn"},
            {"uFillTx", "uFill"},
        )
        exclusive = next((group for group in choices if local in group), {local})
        existing = next((x for x in result if x.tag == child.tag), None)
        clone = (
            _merge_properties(existing, child)
            if local == "defRPr" and existing is not None
            else deepcopy(child)
        )
        for old in list(result):
            if old.tag.rsplit("}", 1)[-1] in exclusive:
                result.remove(old)
        result.append(clone)
    kind = result.tag.rsplit("}", 1)[-1]
    if kind in {"rPr", "defRPr", "endParaRPr"}:
        order = (
            "ln",
            "noFill",
            "solidFill",
            "gradFill",
            "blipFill",
            "pattFill",
            "grpFill",
            "effectLst",
            "effectDag",
            "highlight",
            "uLnTx",
            "uLn",
            "uFillTx",
            "uFill",
            "latin",
            "ea",
            "cs",
            "sym",
            "hlinkClick",
            "hlinkMouseOver",
            "rtl",
            "extLst",
        )
    elif kind == "bodyPr":
        order = (
            "prstTxWarp",
            "noAutofit",
            "normAutofit",
            "spAutoFit",
            "scene3d",
            "sp3d",
            "flatTx",
            "extLst",
        )
    else:
        return result
    for child in sorted(
        list(result),
        key=lambda node: (
            order.index(node.tag.rsplit("}", 1)[-1])
            if node.tag.rsplit("}", 1)[-1] in order
            else len(order)
        ),
    ):
        result.remove(child)
        result.append(child)
    return result


def materialize_text_shape(shape: Any) -> Any:
    """Bake the source placeholder inheritance before moving to a blank layout."""
    clone = deepcopy(shape._element)
    if not shape.has_text_frame:
        return clone
    slide = shape.part.slide
    master = slide.slide_layout.slide_master
    inherited = []
    style_name = "otherStyle"
    if shape.is_placeholder:
        kind = str(shape.placeholder_format.type)
        style_name = "titleStyle" if "TITLE" in kind else "bodyStyle"
        layout_shape = slide.slide_layout.placeholders.get(shape.placeholder_format.idx)
        master_shape = next(
            (
                s
                for s in master.placeholders
                if s.placeholder_format.type == shape.placeholder_format.type
            ),
            None,
        )
        inherited = [s for s in (master_shape, layout_shape) if s is not None and s.has_text_frame]
    body = clone.find("{http://schemas.openxmlformats.org/presentationml/2006/main}txBody")
    body_props = None
    for ancestor in [*inherited, shape]:
        body_props = _merge_properties(body_props, ancestor.text_frame._txBody.bodyPr)
    body_props.tag = body.bodyPr.tag
    body.replace(body.bodyPr, body_props)
    for paragraph in body.p_lst:
        level = int(paragraph.pPr.get("lvl", "0")) if paragraph.pPr is not None else 0
        roots = master._element.xpath(f"./p:txStyles/p:{style_name}/a:lvl{level + 1}pPr")
        effective = deepcopy(roots[0]) if roots else OxmlElement("a:pPr")
        effective.tag = OxmlElement("a:pPr").tag
        for ancestor in [*inherited, shape]:
            frame = ancestor.text_frame._txBody
            defaults = frame.xpath(f"./a:lstStyle/a:lvl{level + 1}pPr")
            if defaults:
                effective = _merge_properties(effective, defaults[0])
            if ancestor is not shape:
                matches = [
                    p
                    for p in frame.p_lst
                    if (int(p.pPr.get("lvl", "0")) if p.pPr is not None else 0) == level
                ]
                if matches:
                    effective = _merge_properties(effective, matches[0].pPr)
        effective = _merge_properties(effective, paragraph.pPr)
        order = (
            "lnSpc",
            "spcBef",
            "spcAft",
            "buClrTx",
            "buClr",
            "buSzTx",
            "buSzPct",
            "buSzPts",
            "buFontTx",
            "buFont",
            "buNone",
            "buAutoNum",
            "buChar",
            "buBlip",
            "tabLst",
            "defRPr",
            "extLst",
        )
        for child in sorted(
            list(effective),
            key=lambda node: (
                order.index(node.tag.rsplit("}", 1)[-1])
                if node.tag.rsplit("}", 1)[-1] in order
                else len(order)
            ),
        ):
            effective.remove(child)
            effective.append(child)
        font_defaults = effective.find(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}defRPr"
        )
        for run in paragraph.xpath("./a:r|./a:fld"):
            props = run.find("{http://schemas.openxmlformats.org/drawingml/2006/main}rPr")
            merged = _merge_properties(font_defaults, props)
            merged.tag = OxmlElement("a:rPr").tag
            if props is not None:
                run.remove(props)
            run.insert(0, merged)
        if paragraph.pPr is not None:
            paragraph.remove(paragraph.pPr)
        effective.set("lvl", str(level))
        paragraph.insert(0, effective)
    for ph in clone.xpath("./p:nvSpPr/p:nvPr/p:ph"):
        ph.getparent().remove(ph)
    return clone


def remove_shape(shape: Any) -> None:
    element = shape._element
    element.getparent().remove(element)


def _remap_relationships(element: Any, source_slide: Any, target_slide: Any) -> None:
    namespace = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    mapping = {}
    for node in element.iter():
        for attr, value in node.attrib.items():
            if not attr.startswith(namespace) or value in mapping:
                continue
            if value not in source_slide.part.rels:
                continue
            relationship = source_slide.part.rels[value]
            target = (
                relationship.target_ref if relationship.is_external else relationship.target_part
            )
            mapping[value] = target_slide.part.relate_to(
                target, relationship.reltype, relationship.is_external
            )
    for node in element.iter():
        for attr, value in list(node.attrib.items()):
            if attr.startswith(namespace) and value in mapping:
                node.set(attr, mapping[value])


def append_element(target_slide: Any, element: Any, source_slide: Any) -> Any:
    clone = deepcopy(element)
    _remap_relationships(clone, source_slide, target_slide)
    next_id = target_slide.shapes._next_shape_id
    for offset, node in enumerate(clone.xpath(".//p:cNvPr")):
        node.set("id", str(next_id + offset))
    target_slide.shapes._spTree.insert_element_before(clone, "p:extLst")
    return target_slide.shapes[-1]


def clone_blank_slide(presentation: Any, source_slide: Any, insert_after: Any) -> Any:
    layouts = source_slide.slide_layout.slide_master.slide_layouts
    blank = min(layouts, key=lambda layout: len(layout.placeholders))
    target = presentation.slides.add_slide(blank)
    for shape in list(target.shapes):
        remove_shape(shape)
    ids = presentation.slides._sldIdLst
    new_id = ids[-1]
    ids.remove(new_id)
    index = next(i for i, slide in enumerate(presentation.slides) if slide is insert_after)
    ids.insert(index + 1, new_id)
    return target
