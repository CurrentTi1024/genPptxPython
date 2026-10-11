from __future__ import annotations

import os
import tempfile
from collections.abc import Mapping
from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from lxml.etree import XMLSyntaxError
from pptx import Presentation
from pptx.exc import PackageNotFoundError

from .config import GeneratorConfig
from .errors import TemplateError
from .flow_renderer import Resources, expand_pages
from .image_renderer import ImageCache, replace_single_images
from .planner import analyze_template
from .text_renderer import TextBudget, _process_shape


def _save_atomically(presentation: Any, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{output_path.stem}.", suffix=".tmp.pptx", dir=output_path.parent
    )
    os.close(descriptor)
    temporary_path = Path(temporary_name)
    try:
        presentation.save(str(temporary_path))
        os.replace(temporary_path, output_path)
    finally:
        temporary_path.unlink(missing_ok=True)


def generate_presentation(
    template_path: str | Path,
    data: Mapping[str, Any],
    output_path: str | Path,
    *,
    config: GeneratorConfig | None = None,
) -> Path:
    """Expand every original template page with a single, atomic generation pass."""
    config = config or GeneratorConfig()
    template_path, output_path = Path(template_path), Path(output_path)
    if not template_path.is_file():
        raise TemplateError(f"template does not exist: {template_path}")
    if template_path.resolve() == output_path.resolve():
        raise TemplateError("output path must differ from template path")
    if not isinstance(data, Mapping):
        raise TemplateError("root payload must be an object")
    if template_path.stat().st_size > config.max_template_bytes:
        raise TemplateError("template file size exceeds limit")
    with template_path.open("rb") as stream:
        template_bytes = stream.read(config.max_template_bytes + 1)
    if len(template_bytes) > config.max_template_bytes:
        raise TemplateError("template file size exceeds limit")
    try:
        with ZipFile(BytesIO(template_bytes)) as archive:
            parts = archive.infolist()
            if (
                len(parts) > config.max_template_parts
                or sum(part.file_size for part in parts) > config.max_template_uncompressed_bytes
            ):
                raise TemplateError("uncompressed template size or part count exceeds limit")
        presentation = Presentation(BytesIO(template_bytes))
    except (BadZipFile, XMLSyntaxError, PackageNotFoundError, KeyError, ValueError) as exc:
        raise TemplateError("template is not a valid PowerPoint package") from exc
    if len(presentation.slides) > config.max_output_slides:
        raise TemplateError("template slide count exceeds limit")
    plans = analyze_template(presentation)
    text_budget = TextBudget(config)
    # Root dynamic text is claimed by the page planner; do not render it twice.
    for plan in plans:
        for shape in list(plan.slide.shapes):
            if shape.shape_id in plan.dynamic_ids:
                continue
            _process_shape(shape, data, text_budget)
    cache = ImageCache({})
    _, images = replace_single_images(presentation, data, config, cache)
    resources = Resources(config, text_budget, images=images)
    expand_pages(presentation, plans, data, resources, cache)
    _save_atomically(presentation, output_path)
    return output_path
