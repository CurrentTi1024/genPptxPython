from __future__ import annotations

import math
import unicodedata
from copy import deepcopy
from dataclasses import dataclass
from typing import Any, Sequence

from pptx.util import Inches, Pt

from .errors import TemplateError
from .models import NormalizedTable, ShapePrototype


@dataclass(frozen=True)
class PreparedTable:
    column_indices: list[int]
    labels: list[str]
    rows: list[list[str]]
    widths: list[int]
    row_heights: list[int]
    font_sizes: list[list[float | None]]


def source_column_index(source_count: int, target_index: int, target_count: int) -> int:
    if source_count <= 1 or target_count <= 1:
        return 0
    if target_index == 0:
        return 0
    if target_index == target_count - 1:
        return source_count - 1
    if source_count == 2:
        return 0
    ratio = target_index / (target_count - 1)
    return min(source_count - 2, max(1, round(ratio * (source_count - 1))))


def source_row_index(source_count: int, data_index: int | None) -> int:
    if data_index is None or source_count <= 1:
        return 0
    return 1 + (data_index % (source_count - 1))


def copy_table_properties(source_table: Any, target_table: Any) -> None:
    target_table._tbl.replace(target_table._tbl.tblPr, deepcopy(source_table._tbl.tblPr))


def copy_cell_style(source_cell: Any, target_cell: Any) -> None:
    if source_cell._tc.tcPr is not None and target_cell._tc.tcPr is not None:
        target_cell._tc.replace(target_cell._tc.tcPr, deepcopy(source_cell._tc.tcPr))
    target_cell.margin_left = source_cell.margin_left
    target_cell.margin_right = source_cell.margin_right
    target_cell.margin_top = source_cell.margin_top
    target_cell.margin_bottom = source_cell.margin_bottom
    target_cell.vertical_anchor = source_cell.vertical_anchor


def set_cell_text(
    target_cell: Any,
    value: str,
    source_cell: Any,
    font_size_points: float | None = None,
) -> None:
    target_cell.text = value
    source_paragraph = source_cell.text_frame.paragraphs[0]
    target_paragraph = target_cell.text_frame.paragraphs[0]
    target_paragraph.alignment = source_paragraph.alignment
    target_paragraph.level = source_paragraph.level
    if source_paragraph.runs and target_paragraph.runs:
        source_r_pr = source_paragraph.runs[0]._r.get_or_add_rPr()
        target_r_pr = target_paragraph.runs[0]._r.get_or_add_rPr()
        target_paragraph.runs[0]._r.replace(target_r_pr, deepcopy(source_r_pr))
        if font_size_points is not None:
            target_paragraph.runs[0].font.size = Pt(font_size_points)


def split_column_indices(
    column_count: int,
    table_width: int,
    min_width_inches: float,
    repeat_leading_columns: int,
) -> list[list[int]]:
    max_columns = int(table_width // Inches(min_width_inches))
    if max_columns < 1:
        raise TemplateError("table width is smaller than the configured minimum column width")
    if column_count <= max_columns:
        return [list(range(column_count))]
    if repeat_leading_columns >= max_columns:
        raise TemplateError(
            "repeatLeadingColumns must leave room for at least one non-repeated column"
        )
    repeat = min(repeat_leading_columns, column_count)
    groups = [list(range(max_columns))]
    cursor = max_columns
    payload_size = max_columns - repeat
    while cursor < column_count:
        groups.append(
            list(range(repeat))
            + list(range(cursor, min(column_count, cursor + payload_size)))
        )
        cursor += payload_size
    return groups


def _content_units(text: str) -> float:
    return max(
        1.0,
        sum(
            1.0 if unicodedata.east_asian_width(character) in {"W", "F", "A"} else 0.55
            for character in text
        ),
    )


def _allocate_widths(total: int, weights: Sequence[float], minimum: int) -> list[int]:
    if minimum * len(weights) > total:
        raise TemplateError("column fragment cannot satisfy the minimum readable width")
    remaining = total - minimum * len(weights)
    weight_sum = sum(weights) or float(len(weights))
    extras = [int(remaining * weight / weight_sum) for weight in weights]
    for index in range(remaining - sum(extras)):
        extras[index % len(extras)] += 1
    widths = [minimum + extra for extra in extras]
    return widths


def column_widths(
    source_table: Any,
    normalized: NormalizedTable,
    column_indices: Sequence[int],
) -> list[int]:
    total_width = sum(column.width for column in source_table.columns)
    minimum = Inches(normalized.options.min_column_width_inches or 0.65)
    if normalized.options.column_width == "equal":
        weights = [1.0] * len(column_indices)
    elif normalized.options.column_width == "content":
        weights = []
        for index in column_indices:
            values = [normalized.labels[index], *(row[index] for row in normalized.rows)]
            weights.append(min(40.0, max(_content_units(value) for value in values)))
    else:
        source_widths = [column.width for column in source_table.columns]
        weights = [
            float(
                source_widths[
                    source_column_index(
                        len(source_widths), target_index, len(column_indices)
                    )
                ]
            )
            for target_index in range(len(column_indices))
        ]
    return _allocate_widths(total_width, weights, minimum)


def _font_size(source_cell: Any) -> float:
    paragraph = source_cell.text_frame.paragraphs[0]
    if paragraph.runs and paragraph.runs[0].font.size:
        return paragraph.runs[0].font.size.pt
    return 12.0


def _line_count(text: str, width_points: float, font_points: float) -> int:
    capacity = max(1.0, width_points / max(1.0, font_points))
    return max(
        1,
        sum(max(1, math.ceil(_content_units(line) / capacity)) for line in text.split("\n")),
    )


def _truncate(text: str, width_points: float, font_points: float, lines: int) -> str:
    capacity = max(1, int(width_points / max(1.0, font_points) * lines / 0.7))
    if len(text) <= capacity:
        return text
    return text[: max(0, capacity - 1)].rstrip() + "…"


def prepare_table(
    source_table: Any,
    normalized: NormalizedTable,
    column_indices: list[int],
) -> PreparedTable:
    labels = [normalized.labels[index] for index in column_indices]
    rows = [[row[index] for index in column_indices] for row in normalized.rows]
    widths = column_widths(source_table, normalized, column_indices)
    all_rows = [labels, *rows]
    heights: list[int] = []
    font_sizes: list[list[float | None]] = []
    max_height = Inches(normalized.options.max_row_height_inches or 1.4)
    min_font = normalized.options.min_font_size_points or 8.0

    for row_index, values in enumerate(all_rows):
        source_index = source_row_index(
            len(source_table.rows), None if row_index == 0 else row_index - 1
        )
        required_height = 0
        row_fonts: list[float | None] = []
        metrics: list[tuple[float, float, int]] = []
        for target_index, (value, width) in enumerate(zip(values, widths)):
            source_col = source_column_index(
                len(source_table.columns), target_index, len(column_indices)
            )
            cell = source_table.cell(source_index, source_col)
            size = _font_size(cell)
            horizontal = (cell.margin_left or 0) + (cell.margin_right or 0)
            vertical = (cell.margin_top or 0) + (cell.margin_bottom or 0)
            width_points = max(1.0, (width - horizontal) / 12700)
            lines = _line_count(value, width_points, size)
            cell_height = vertical + Pt(size * 1.22 * lines)
            required_height = max(required_height, cell_height)
            metrics.append((width_points, size, vertical))
            row_fonts.append(None)
        template_height = source_table.rows[source_index].height
        required_height = max(template_height, required_height)

        overflow = normalized.options.overflow
        if required_height > max_height and overflow == "error":
            raise TemplateError(f"table row {row_index} exceeds maxRowHeight")
        if required_height > max_height and overflow in {"shrink", "truncate"}:
            for index, (value, metric) in enumerate(zip(values, metrics)):
                width_points, base_size, vertical = metric
                chosen_size = base_size
                if overflow == "shrink":
                    chosen_size = max(min_font, base_size * max_height / required_height)
                    row_fonts[index] = chosen_size
                usable_height = max(1.0, (max_height - vertical) / 12700)
                allowed_lines = max(1, int(usable_height / (chosen_size * 1.22)))
                all_rows[row_index][index] = _truncate(
                    value, width_points, chosen_size, allowed_lines
                )
            required_height = max_height
        heights.append(int(required_height))
        font_sizes.append(row_fonts)
    return PreparedTable(column_indices, all_rows[0], all_rows[1:], widths, heights, font_sizes)


def add_table_fragment(
    slide: Any,
    prototype: ShapePrototype,
    prepared: PreparedTable,
    left: int,
    top: int,
    row_offset: int,
    row_count: int,
    generated_name: str,
) -> tuple[Any, int]:
    source_table = prototype.shape.table
    row_heights = [prepared.row_heights[0]] + prepared.row_heights[
        1 + row_offset : 1 + row_offset + row_count
    ]
    table_height = sum(row_heights)
    table_shape = slide.shapes.add_table(
        1 + row_count,
        len(prepared.labels),
        left,
        top,
        prototype.shape.width,
        table_height,
    )
    table_shape.name = generated_name
    target_table = table_shape.table
    copy_table_properties(source_table, target_table)
    for index, width in enumerate(prepared.widths):
        target_table.columns[index].width = width
    for index, height in enumerate(row_heights):
        target_table.rows[index].height = height

    values = [prepared.labels, *prepared.rows[row_offset : row_offset + row_count]]
    font_rows = [
        prepared.font_sizes[0],
        *prepared.font_sizes[1 + row_offset : 1 + row_offset + row_count],
    ]
    for target_row_index, (row_values, row_fonts) in enumerate(zip(values, font_rows)):
        global_data_index = None if target_row_index == 0 else row_offset + target_row_index - 1
        source_row = source_row_index(len(source_table.rows), global_data_index)
        for target_col_index, (value, font_size) in enumerate(zip(row_values, row_fonts)):
            source_col = source_column_index(
                len(source_table.columns), target_col_index, len(prepared.labels)
            )
            source_cell = source_table.cell(source_row, source_col)
            target_cell = target_table.cell(target_row_index, target_col_index)
            copy_cell_style(source_cell, target_cell)
            set_cell_text(target_cell, value, source_cell, font_size)
    return table_shape, table_height
