from __future__ import annotations

import json
import unittest
from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
from pathlib import Path
from tempfile import TemporaryDirectory
from threading import Event
from types import SimpleNamespace
from unittest.mock import patch

from fastapi.testclient import TestClient
from pptx import Presentation

from genppt import GeneratorConfig
from genppt.core.settings import Settings
from genppt.main import create_app

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures"
TEMPLATE = FIXTURES / "test-multipage.pptx"
ENDPOINT = "/api/v1/presentations"


class ApiTest(unittest.TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.settings = Settings(
            output_dir=Path(self.directory.name) / "output",
            template_roots=(FIXTURES,),
            generator=GeneratorConfig(allowed_image_roots=(str(FIXTURES / "images"),)),
        )
        self.app = create_app(self.settings)
        self.client = TestClient(self.app)
        self.addCleanup(self.client.close)

    def test_json_path_route_returns_path_and_generates_ppt(self):
        response = self.client.post(ENDPOINT, json={"template": str(TEMPLATE), "data": "{}"})
        self.assertEqual(response.status_code, 200)
        self.assertIsInstance(response.json(), str)
        self.assertEqual(len(Presentation(response.json()).slides), 3)

    def test_multipart_route_real_template_complete_data(self):
        from examples.run_multipage_demo import demo_payload

        response = self.client.post(
            ENDPOINT,
            files={"template": ("report.pptx", TEMPLATE.read_bytes())},
            data={"data": json.dumps(demo_payload(), ensure_ascii=False)},
        )
        self.assertEqual(response.status_code, 200)
        result = Presentation(response.json())
        tables = [s.table for slide in result.slides for s in slide.shapes if s.has_table]
        self.assertEqual(len(result.slides), 6)
        self.assertEqual(len(tables), 7)
        self.assertEqual(sum(len(t.rows) - 1 for t in tables), 31)
        self.assertEqual(sum(s.shape_type == 13 for sl in result.slides for s in sl.shapes), 4)

    def test_upload_filename_is_not_used_as_output_path(self):
        response = self.client.post(
            ENDPOINT,
            files={"template": ("../../outside.pptx", TEMPLATE.read_bytes())},
            data={"data": "{}"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Path(response.json()).parent, self.settings.output_dir)
        self.assertNotEqual(Path(response.json()).name, "outside.pptx")

    def test_path_root_traversal_and_symlink_are_denied(self):
        outside = Path(self.directory.name) / "outside.pptx"
        outside.write_bytes(TEMPLATE.read_bytes())
        for path in (str(outside), str(FIXTURES / "../../README.md")):
            response = self.client.post(ENDPOINT, json={"template": path, "data": "{}"})
            self.assertEqual(response.status_code, 403)
        allowed = Path(self.directory.name) / "allowed"
        allowed.mkdir()
        link = allowed / "link.pptx"
        link.symlink_to(outside)
        with TestClient(create_app(replace(self.settings, template_roots=(allowed,)))) as client:
            response = client.post(ENDPOINT, json={"template": str(link), "data": "{}"})
            self.assertEqual(response.status_code, 403)

    def test_unknown_missing_or_wrong_fields_are_rejected(self):
        for body in (
            {"template": str(TEMPLATE)},
            {"template": str(TEMPLATE), "data": {}},
            {"template": str(TEMPLATE), "data": "{}", "output": "/tmp/other"},
            [str(TEMPLATE), "{}"],
        ):
            with self.subTest(body=body):
                self.assertEqual(self.client.post(ENDPOINT, json=body).status_code, 422)

    def test_invalid_inner_json_and_template_are_422(self):
        for data in ("[]", "{broken}", '{"a":NaN}', '{"a":"\\u0000"}'):
            response = self.client.post(ENDPOINT, json={"template": str(TEMPLATE), "data": data})
            self.assertEqual(response.status_code, 422)
        response = self.client.post(
            ENDPOINT, files={"template": ("broken.pptx", b"bad zip")}, data={"data": "{}"}
        )
        self.assertEqual(response.status_code, 422)
        self.assertFalse(self.settings.output_dir.exists())

    def test_unsupported_content_type_is_415(self):
        self.assertEqual(
            self.client.post(ENDPOINT, data={"template": "x", "data": "{}"}).status_code, 415
        )

    def test_upload_wrong_type_and_extra_field_are_rejected(self):
        for fields in ({"data": "{}", "extra": "x"}, {"wrong": "{}"}):
            response = self.client.post(
                ENDPOINT, files={"template": ("report.pptx", TEMPLATE.read_bytes())}, data=fields
            )
            self.assertIn(response.status_code, (400, 422))
        response = self.client.post(
            ENDPOINT, files={"data": ("data.txt", b"{}")}, data={"template": "report.pptx"}
        )
        self.assertEqual(response.status_code, 422)

    def test_template_upload_and_data_size_limits(self):
        config = replace(self.settings.generator, max_template_bytes=10)
        with TestClient(create_app(replace(self.settings, generator=config))) as client:
            response = client.post(
                ENDPOINT,
                files={"template": ("report.pptx", TEMPLATE.read_bytes())},
                data={"data": "{}"},
            )
            self.assertEqual(response.status_code, 413)
        with TestClient(create_app(replace(self.settings, max_data_bytes=8))) as client:
            response = client.post(
                ENDPOINT, json={"template": str(TEMPLATE), "data": '{"x":"too big"}'}
            )
            self.assertEqual(response.status_code, 422)
            response = client.post(
                ENDPOINT,
                files={"template": ("report.pptx", TEMPLATE.read_bytes())},
                data={"data": '{"x":"too big"}'},
            )
            self.assertEqual(response.status_code, 400)

    def test_total_body_limit_including_chunked_requests(self):
        settings = replace(
            self.settings, max_data_bytes=1, generator=GeneratorConfig(max_template_bytes=1)
        )
        with TestClient(create_app(settings)) as client:
            response = client.post(
                ENDPOINT,
                content=b"x" * (settings.max_request_bytes + 1),
                headers={"content-type": "application/json"},
            )
            self.assertEqual(response.status_code, 413)
            response = client.post(
                ENDPOINT,
                content=iter((b"x" * (settings.max_request_bytes + 1),)),
                headers={"content-type": "application/json"},
            )
            self.assertEqual(response.status_code, 413)

    def test_worker_failure_has_sanitized_response_and_no_file(self):
        with (
            TestClient(self.app, raise_server_exceptions=False) as client,
            patch(
                "pptx.presentation.Presentation.save", side_effect=OSError("secret internal path")
            ),
            self.assertLogs("genppt.main", level="ERROR"),
        ):
            response = client.post(ENDPOINT, json={"template": str(TEMPLATE), "data": "{}"})
        self.assertEqual(response.status_code, 500)
        self.assertEqual(response.json(), {"detail": "generation failed"})
        self.assertEqual(list(self.settings.output_dir.iterdir()), [])

    def test_upload_stream_is_closed_after_worker_finishes(self):
        seen = []

        def generate(template, data):
            self.assertFalse(template.closed)
            seen.append(template)
            return "/tmp/generated.pptx"

        self.app.state.presentation_service = SimpleNamespace(
            settings=self.settings, generate=generate
        )
        response = self.client.post(
            ENDPOINT,
            files={"template": ("report.pptx", TEMPLATE.read_bytes())},
            data={"data": "{}"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(seen[0].closed)

    def test_admission_rejects_excess_jobs_and_recovers(self):
        started, release = Event(), Event()
        app = create_app(replace(self.settings, max_concurrent_jobs=1))

        def generate(template, data):
            started.set()
            if not release.wait(5):
                raise AssertionError("worker was not released")
            return "/tmp/generated.pptx"

        app.state.presentation_service = SimpleNamespace(settings=self.settings, generate=generate)
        with TestClient(app) as client, ThreadPoolExecutor(max_workers=1) as pool:
            first = pool.submit(
                client.post, ENDPOINT, json={"template": str(TEMPLATE), "data": "{}"}
            )
            try:
                self.assertTrue(started.wait(5))
                second = client.post(ENDPOINT, json={"template": str(TEMPLATE), "data": "{}"})
                self.assertEqual(second.status_code, 503)
                self.assertEqual(second.headers["Retry-After"], "1")
                self.assertEqual(client.get("/health").status_code, 200)
            finally:
                release.set()
            self.assertEqual(first.result(timeout=5).status_code, 200)
            self.assertEqual(
                client.post(ENDPOINT, json={"template": str(TEMPLATE), "data": "{}"}).status_code,
                200,
            )

    def test_openapi_documents_both_inputs_and_health(self):
        schema = self.client.get("/openapi.json").json()
        content = schema["paths"][ENDPOINT]["post"]["requestBody"]["content"]
        self.assertEqual(set(content), {"application/json", "multipart/form-data"})
        for spec in content.values():
            self.assertEqual(set(spec["schema"]["required"]), {"template", "data"})
        self.assertEqual(self.client.get("/health").json(), {"status": "ok"})

    def test_admission_policy_also_applies_under_mount_prefix(self):
        from genppt.api.middleware import GenerationAdmissionMiddleware

        async def unused_app(scope, receive, send):
            raise AssertionError("request should have been rejected before reaching the app")

        middleware = GenerationAdmissionMiddleware(unused_app, max_jobs=1)
        middleware.slots.acquire()
        client = TestClient(middleware, root_path="/ppt")
        self.addCleanup(client.close)
        response = client.post("/ppt/api/v1/presentations", json={})
        self.assertEqual(response.status_code, 503)
        middleware.slots.release()

    def test_upload_stream_is_closed_after_generation_failure(self):
        from genppt import TemplateError

        seen = []

        def generate(template, data):
            seen.append(template)
            raise TemplateError("invalid template")

        self.app.state.presentation_service = SimpleNamespace(
            settings=self.settings, generate=generate
        )
        response = self.client.post(
            ENDPOINT,
            files={"template": ("report.pptx", TEMPLATE.read_bytes())},
            data={"data": "{}"},
        )
        self.assertEqual(response.status_code, 422)
        self.assertTrue(seen[0].closed)


if __name__ == "__main__":
    unittest.main()
