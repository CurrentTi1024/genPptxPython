from __future__ import annotations

import json
import unittest
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory

from pptx import Presentation
from pptx.util import Inches

from genppt.config import GeneratorConfig
from genppt.data import normalize_table
from genppt.errors import TemplateError
from genppt.service import generate_presentation


ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "tests" / "fixtures"
CONFIG = GeneratorConfig()


class NormalizeTableTest(unittest.TestCase):
    def test_matrix_rows(self) -> None:
        result = normalize_table({"columns": ["A", "B"], "rows": [[1, 2]]}, CONFIG)
        self.assertEqual(result.labels, ["A", "B"])
        self.assertEqual(result.rows, [["1", "2"]])

    def test_object_rows(self) -> None:
        result = normalize_table(
            {"columns": ["name", "count"], "rows": [{"name": "华东", "count": 3}]},
            CONFIG,
        )
        self.assertEqual(result.rows, [["华东", "3"]])

    def test_row_width_mismatch_is_rejected(self) -> None:
        with self.assertRaises(TemplateError):
            normalize_table({"columns": ["A", "B"], "rows": [[1]]}, CONFIG)

    def test_non_finite_number_is_rejected(self) -> None:
        with self.assertRaises(TemplateError):
            normalize_table({"columns": ["A"], "rows": [[float("nan")]]}, CONFIG)

    def test_resource_limit_is_enforced(self) -> None:
        with self.assertRaises(TemplateError):
            normalize_table(
                {"columns": ["A", "B"], "rows": []},
                GeneratorConfig(max_columns=1),
            )

    def test_column_format_and_null_value(self) -> None:
        result = normalize_table(
            {
                "columns": [
                    {"key": "amount", "label": "金额", "format": ",.2f"},
                    {"key": "day", "label": "日期", "format": "%Y/%m/%d"},
                    {"key": "note", "label": "备注", "nullValue": "-"},
                ],
                "rows": [
                    {"amount": Decimal("1234.5"), "day": "2026-09-20", "note": None}
                ],
                "_table": {"nullValue": "N/A"},
            },
            CONFIG,
        )
        self.assertEqual(result.rows, [["1,234.50", "2026/09/20", "-"]])


def _add_block(
    slide,
    item: str,
    block: str,
    top: float,
    left: float = 0.7,
    width: float = 6.0,
) -> None:
    title = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(0.3)
    )
    title.name = f"<foreach item={item} block={block}>{{{{title}}}}</foreach>"
    title.text = "{{title}}{{_page.continuationSuffix}}"
    table_shape = slide.shapes.add_table(
        2, 2, Inches(left), Inches(top + 0.4), Inches(width), Inches(0.8)
    )
    table_shape.name = f"<foreach item={item} block={block}>{{{{name}}}}</foreach>"
    table_shape.table.cell(0, 0).text = "A"
    table_shape.table.cell(0, 1).text = "B"


def _multi_block_template(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _add_block(slide, "primary", "primaryBlock", 1.0)
    _add_block(slide, "secondary", "secondaryBlock", 4.0)
    presentation.save(path)


class GeneratePresentationTest(unittest.TestCase):
    def test_dynamic_cell_value_is_not_reinterpreted_as_root_expression(self) -> None:
        data = {
            "tables": [
                {
                    "title": "字面量",
                    "columns": ["内容"],
                    "rows": [["保留 {{rootValue}} 原文"]],
                }
            ],
            "rootValue": "不应替换",
        }
        with TemporaryDirectory() as directory:
            output = Path(directory) / "generated.pptx"
            generate_presentation(FIXTURES / "test2.pptx", data, output)
            presentation = Presentation(output)
            table = next(
                shape.table
                for slide in presentation.slides
                for shape in slide.shapes
                if shape.has_table
            )
            self.assertEqual(table.cell(1, 0).text, "保留 {{rootValue}} 原文")

    def test_test_template_generates_dynamic_tables(self) -> None:
        data = json.loads((ROOT / "examples/test_data.json").read_text(encoding="utf-8"))
        with TemporaryDirectory() as directory:
            output = Path(directory) / "generated.pptx"
            generate_presentation(FIXTURES / "test2.pptx", data, output)
            presentation = Presentation(output)
            tables = [
                shape
                for slide in presentation.slides
                for shape in slide.shapes
                if shape.has_table
            ]
            self.assertEqual(len(presentation.slides), 2)
            self.assertEqual(len(tables), 3)
            self.assertEqual(len(tables[0].table.columns), 3)
            self.assertEqual(len(tables[0].table.rows), 3)
            self.assertEqual(tables[0].table.cell(0, 0).text, "区域")
            self.assertEqual(tables[0].table.cell(1, 0).text, "华东区")
            self.assertTrue(
                any(
                    "分类：资料不完整" in shape.text
                    for shape in presentation.slides[0].shapes
                    if getattr(shape, "has_text_frame", False)
                )
            )
            self.assertTrue(
                all(
                    "<foreach" not in shape.name
                    for slide in presentation.slides
                    for shape in slide.shapes
                )
            )

    def test_columns_can_grow_beyond_template_width(self) -> None:
        data = {
            "tables": [
                {
                    "title": "九列表格",
                    "tableName": "nine_columns",
                    "columns": [f"C{index}" for index in range(1, 10)],
                    "rows": [[f"V{index}" for index in range(1, 10)]],
                }
            ]
        }
        with TemporaryDirectory() as directory:
            output = Path(directory) / "generated.pptx"
            generate_presentation(FIXTURES / "test2.pptx", data, output)
            presentation = Presentation(output)
            tables = [shape for shape in presentation.slides[0].shapes if shape.has_table]
            self.assertEqual(len(tables), 1)
            self.assertEqual(len(tables[0].table.columns), 9)
            self.assertEqual(len(tables[0].table.rows), 2)
            self.assertEqual(tables[0].table.cell(1, 8).text, "V9")

    def test_two_short_tables_share_a_compact_page(self) -> None:
        data = {
            "tables": [
                {"title": "A", "tableName": "a", "columns": ["列"], "rows": [[1]]},
                {"title": "B", "tableName": "b", "columns": ["列"], "rows": [[2]]},
            ]
        }
        with TemporaryDirectory() as directory:
            output = Path(directory) / "generated.pptx"
            generate_presentation(FIXTURES / "test2.pptx", data, output)
            presentation = Presentation(output)
            tables = [shape for shape in presentation.slides[0].shapes if shape.has_table]
            self.assertEqual(len(presentation.slides), 1)
            self.assertTrue(tables[0].name.startswith("a__reasonTable_i1"))
            self.assertTrue(tables[1].name.startswith("b__reasonTable_i2"))

    def test_next_table_moves_whole_to_new_page_when_remainder_is_too_small(self) -> None:
        data = {
            "tables": [
                {
                    "title": "第一张",
                    "tableName": "first",
                    "columns": ["序号", "值"],
                    "rows": [[index, index] for index in range(1, 8)],
                },
                {
                    "title": "第二张",
                    "tableName": "second",
                    "columns": ["序号", "值"],
                    "rows": [[index, index] for index in range(1, 4)],
                },
            ]
        }
        with TemporaryDirectory() as directory:
            output = Path(directory) / "generated.pptx"
            generate_presentation(FIXTURES / "test2.pptx", data, output)
            presentation = Presentation(output)
            tables_by_slide = [
                [shape for shape in slide.shapes if shape.has_table]
                for slide in presentation.slides
            ]
            self.assertEqual(len(presentation.slides), 2)
            self.assertTrue(tables_by_slide[0][0].name.startswith("first__reasonTable_i1"))
            self.assertTrue(tables_by_slide[1][0].name.startswith("second__reasonTable_i2"))
            self.assertEqual(len(tables_by_slide[1][0].table.rows), 4)

    def test_empty_rows_remove_the_entire_repeated_block(self) -> None:
        data = {
            "tables": [
                {"title": "不应显示", "columns": ["A"], "rows": []},
                {"title": "显示", "columns": ["A"], "rows": [[1]]},
            ]
        }
        with TemporaryDirectory() as directory:
            output = Path(directory) / "generated.pptx"
            generate_presentation(FIXTURES / "test2.pptx", data, output)
            presentation = Presentation(output)
            texts = [
                shape.text
                for shape in presentation.slides[0].shapes
                if getattr(shape, "has_text_frame", False)
            ]
            self.assertFalse(any("不应显示" in text for text in texts))
            self.assertTrue(any("显示" in text for text in texts))

    def test_empty_outer_array_removes_all_placeholders(self) -> None:
        with TemporaryDirectory() as directory:
            output = Path(directory) / "generated.pptx"
            generate_presentation(FIXTURES / "test2.pptx", {"tables": []}, output)
            presentation = Presentation(output)
            self.assertFalse(
                any(shape.has_table for shape in presentation.slides[0].shapes)
            )
            self.assertTrue(
                all(
                    "<foreach" not in shape.name
                    for shape in presentation.slides[0].shapes
                )
            )

    def test_long_text_increases_row_height(self) -> None:
        data = {
            "tables": [
                {
                    "columns": ["说明"],
                    "rows": [
                        [
                            "这是一段需要自动换行"
                            "并增加表格行高的很长中文内容"
                            * 8
                        ]
                    ],
                }
            ]
        }
        with TemporaryDirectory() as directory:
            output = Path(directory) / "generated.pptx"
            generate_presentation(FIXTURES / "test2.pptx", data, output)
            generated = Presentation(output)
            source = Presentation(FIXTURES / "test2.pptx")
            generated_table = next(
                shape.table
                for shape in generated.slides[0].shapes
                if shape.has_table
            )
            source_table = next(
                shape.table for shape in source.slides[0].shapes if shape.has_table
            )
            self.assertGreater(generated_table.rows[1].height, source_table.rows[1].height)

    def test_truncate_overflow_caps_row_and_adds_ellipsis(self) -> None:
        data = {
            "tables": [
                {
                    "columns": ["说明"],
                    "rows": [["非常长的内容" * 80]],
                    "_table": {"overflow": "truncate", "maxRowHeight": 0.5},
                }
            ]
        }
        with TemporaryDirectory() as directory:
            output = Path(directory) / "generated.pptx"
            generate_presentation(FIXTURES / "test2.pptx", data, output)
            presentation = Presentation(output)
            table = next(
                shape.table
                for shape in presentation.slides[0].shapes
                if shape.has_table
            )
            self.assertLessEqual(table.rows[1].height, Inches(0.5))
            self.assertTrue(table.cell(1, 0).text.endswith("…"))

    def test_error_overflow_rejects_excessive_row(self) -> None:
        data = {
            "tables": [
                {
                    "columns": ["说明"],
                    "rows": [["过长内容" * 100]],
                    "_table": {"overflow": "error", "maxRowHeight": 0.4},
                }
            ]
        }
        with TemporaryDirectory() as directory:
            with self.assertRaisesRegex(TemplateError, "exceeds maxRowHeight"):
                generate_presentation(
                    FIXTURES / "test2.pptx",
                    data,
                    Path(directory) / "generated.pptx",
                )

    def test_equal_and_content_column_width_strategies(self) -> None:
        with TemporaryDirectory() as directory:
            equal_output = Path(directory) / "equal.pptx"
            content_output = Path(directory) / "content.pptx"
            base = {
                "columns": ["短", "很长的说明字段"],
                "rows": [["1", "一段明显更长的内容" * 4]],
            }
            generate_presentation(
                FIXTURES / "test2.pptx",
                {"tables": [{**base, "_table": {"columnWidth": "equal"}}]},
                equal_output,
            )
            generate_presentation(
                FIXTURES / "test2.pptx",
                {"tables": [{**base, "_table": {"columnWidth": "content"}}]},
                content_output,
            )
            equal_table = next(
                shape.table
                for shape in Presentation(equal_output).slides[0].shapes
                if shape.has_table
            )
            content_table = next(
                shape.table
                for shape in Presentation(content_output).slides[0].shapes
                if shape.has_table
            )
            self.assertLessEqual(
                abs(equal_table.columns[0].width - equal_table.columns[1].width), 1
            )
            self.assertGreater(
                content_table.columns[1].width, content_table.columns[0].width
            )

    def test_wide_table_is_split_into_readable_column_fragments(self) -> None:
        data = {
            "tables": [
                {
                    "columns": [f"C{index}" for index in range(15)],
                    "rows": [[index for index in range(15)]],
                    "_table": {"minColumnWidth": 1.0, "repeatLeadingColumns": 1},
                }
            ]
        }
        with TemporaryDirectory() as directory:
            output = Path(directory) / "generated.pptx"
            generate_presentation(FIXTURES / "test2.pptx", data, output)
            presentation = Presentation(output)
            tables = [
                shape
                for slide in presentation.slides
                for shape in slide.shapes
                if shape.has_table
            ]
            self.assertGreater(len(tables), 1)
            self.assertTrue(
                all(
                    column.width >= Inches(1.0)
                    for table in tables
                    for column in table.table.columns
                )
            )
            self.assertEqual(tables[1].table.cell(0, 0).text, "C0")
            names = [table.name for table in tables]
            self.assertEqual(len(names), len(set(names)))

    def test_multiple_blocks_use_independent_regions(self) -> None:
        with TemporaryDirectory() as directory:
            template = Path(directory) / "template.pptx"
            output = Path(directory) / "generated.pptx"
            _multi_block_template(template)
            generate_presentation(
                template,
                {
                    "primary": [
                        {"title": "上区", "name": "top", "columns": ["A"], "rows": [[1]]}
                    ],
                    "secondary": [
                        {"title": "下区", "name": "bottom", "columns": ["B"], "rows": [[2]]}
                    ],
                },
                output,
            )
            presentation = Presentation(output)
            tables = [shape for shape in presentation.slides[0].shapes if shape.has_table]
            self.assertEqual(len(tables), 2)
            self.assertLess(tables[0].top, tables[1].top)

    def test_side_by_side_blocks_have_independent_regions(self) -> None:
        with TemporaryDirectory() as directory:
            template = Path(directory) / "template.pptx"
            output = Path(directory) / "generated.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            _add_block(slide, "leftTables", "leftBlock", 1.0, 0.5, 5.5)
            _add_block(slide, "rightTables", "rightBlock", 1.0, 7.0, 5.5)
            presentation.save(template)
            generate_presentation(
                template,
                {
                    "leftTables": [
                        {"title": "左", "name": "left", "columns": ["A"], "rows": [[1]]}
                    ],
                    "rightTables": [
                        {"title": "右", "name": "right", "columns": ["B"], "rows": [[2]]}
                    ],
                },
                output,
            )
            tables = [
                shape
                for shape in Presentation(output).slides[0].shapes
                if shape.has_table
            ]
            self.assertEqual(len(tables), 2)
            self.assertLess(tables[0].left, tables[1].left)

    def test_continuation_title_uses_system_suffix(self) -> None:
        with TemporaryDirectory() as directory:
            template = Path(directory) / "template.pptx"
            output = Path(directory) / "generated.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            _add_block(slide, "tables", "reason", 1.0)
            presentation.save(template)
            generate_presentation(
                template,
                {
                    "tables": [
                        {
                            "title": "长表",
                            "name": "long",
                            "columns": ["A", "B"],
                            "rows": [[index, "x"] for index in range(50)],
                        }
                    ]
                },
                output,
            )
            generated = Presentation(output)
            continuation_texts = [
                shape.text
                for slide in list(generated.slides)[1:]
                for shape in slide.shapes
                if getattr(shape, "has_text_frame", False)
            ]
            self.assertTrue(any("长表（续）" in text for text in continuation_texts))

    def test_table_larger_than_fresh_page_is_split_with_repeated_header(self) -> None:
        data = {
            "tables": [
                {
                    "title": "长表",
                    "tableName": "long_table",
                    "columns": ["序号", "值"],
                    "rows": [[index, f"V{index}"] for index in range(1, 26)],
                }
            ]
        }
        with TemporaryDirectory() as directory:
            output = Path(directory) / "generated.pptx"
            generate_presentation(FIXTURES / "test2.pptx", data, output)
            presentation = Presentation(output)
            tables = [
                shape
                for slide in presentation.slides
                for shape in slide.shapes
                if shape.has_table
            ]
            self.assertGreater(len(tables), 1)
            self.assertTrue(all(table.table.cell(0, 0).text == "序号" for table in tables))
            self.assertEqual(sum(len(table.table.rows) - 1 for table in tables), 25)

    def test_multiple_template_slides_are_rejected_by_default(self) -> None:
        data = {"tables": [{"columns": ["A"], "rows": [[1]]}]}
        with TemporaryDirectory() as directory:
            with self.assertRaises(TemplateError):
                generate_presentation(FIXTURES / "test.pptx", data, Path(directory) / "out.pptx")

    def test_template_cannot_be_overwritten(self) -> None:
        with self.assertRaises(TemplateError):
            generate_presentation(
                FIXTURES / "test2.pptx", {"tables": []}, FIXTURES / "test2.pptx"
            )

    def test_image_foreach_is_routed_to_image_generator(self) -> None:
        with TemporaryDirectory() as directory:
            output = Path(directory) / "out.pptx"
            generate_presentation(
                FIXTURES / "test-img.pptx",
                {"imgs": [], "imageGroups": []},
                output,
            )
            presentation = Presentation(output)
            self.assertFalse(
                any(
                    shape.shape_type == 13
                    for slide in presentation.slides
                    for shape in slide.shapes
                )
            )

    def test_existing_output_survives_validation_failure(self) -> None:
        with TemporaryDirectory() as directory:
            output = Path(directory) / "existing.pptx"
            output.write_bytes(b"existing-output")
            with self.assertRaises(TemplateError):
                generate_presentation(
                    FIXTURES / "test2.pptx",
                    {"tables": [{"columns": ["A", "B"], "rows": [[1]]}]},
                    output,
                )
            self.assertEqual(output.read_bytes(), b"existing-output")


if __name__ == "__main__":
    unittest.main()
