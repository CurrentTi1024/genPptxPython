from __future__ import annotations

import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from pptx import Presentation

from genppt import GeneratorConfig, TemplateError, generate_presentation
from genppt.directives import parse_foreach, parse_single_image


ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "tests" / "fixtures"
IMAGES = FIXTURES / "images"
TEMPLATE = FIXTURES / "test-img.pptx"


def _picture_shapes(presentation: Presentation):
    return [
        shape
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.shape_type == 13
    ]


def _single_image_template(path: Path, name: str) -> tuple[int, int]:
    presentation = Presentation(TEMPLATE)
    picture = next(
        shape for shape in presentation.slides[0].shapes if shape.shape_type == 13
    )
    size = picture.width, picture.height
    picture.name = name
    presentation.save(path)
    return size


def _loop_template(path: Path, name: str) -> None:
    presentation = Presentation(TEMPLATE)
    picture = next(
        shape for shape in presentation.slides[0].shapes if shape.shape_type == 13
    )
    picture.name = name
    presentation.save(path)


class DirectiveTest(unittest.TestCase):
    def test_rows_and_cols_are_optional(self) -> None:
        directive = parse_foreach("<foreach item=imgs></foreach>(x)", 1)
        self.assertIsNotNone(directive)
        self.assertIsNone(directive.rows)
        self.assertIsNone(directive.cols)
        self.assertEqual(directive.image_mode, "x")

    def test_rows_or_cols_can_be_set_independently(self) -> None:
        rows = parse_foreach("<foreach item=imgs rows=2></foreach>", 1)
        cols = parse_foreach("<foreach item=imgs cols=3></foreach>", 1)
        both = parse_foreach(
            "<foreach item=imgs rows=2 cols=3></foreach>(x, y)", 1
        )
        self.assertEqual((rows.rows, rows.cols), (2, None))
        self.assertEqual((cols.rows, cols.cols), (None, 3))
        self.assertEqual((both.rows, both.cols, both.image_mode), (2, 3, "x,y"))

    def test_invalid_grid_dimension_is_rejected(self) -> None:
        with self.assertRaisesRegex(TemplateError, "positive integer"):
            parse_foreach("<foreach item=imgs rows=0></foreach>", 1)

    def test_single_image_default_mode_is_x(self) -> None:
        self.assertEqual(parse_single_image("{{imagePath}}"), ("imagePath", "x"))


class SingleImageTest(unittest.TestCase):
    def _generate(self, directive: str, image: Path | str) -> tuple[Presentation, tuple[int, int]]:
        directory = TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        template = Path(directory.name) / "template.pptx"
        output = Path(directory.name) / "output.pptx"
        original_size = _single_image_template(template, directive)
        generate_presentation(
            template,
            {"image": str(image), "imageGroups": []},
            output,
        )
        return Presentation(output), original_size

    def test_default_and_x_keep_width(self) -> None:
        for directive in ("{{image}}", "{{image}}(x)"):
            with self.subTest(directive=directive):
                presentation, original = self._generate(directive, IMAGES / "wide.png")
                picture = _picture_shapes(presentation)[0]
                self.assertEqual(picture.width, original[0])
                self.assertAlmostEqual(picture.width / picture.height, 3694 / 925, places=2)

    def test_y_keeps_height(self) -> None:
        presentation, original = self._generate("{{image}}(y)", IMAGES / "wide.png")
        picture = _picture_shapes(presentation)[0]
        self.assertEqual(picture.height, original[1])
        self.assertAlmostEqual(picture.width / picture.height, 3694 / 925, places=2)

    def test_xy_keeps_both_dimensions(self) -> None:
        presentation, original = self._generate("{{image}}(x,y)", IMAGES / "wide.png")
        picture = _picture_shapes(presentation)[0]
        self.assertEqual((picture.width, picture.height), original)

    def test_missing_empty_and_corrupt_images_remove_placeholder(self) -> None:
        with TemporaryDirectory() as directory:
            corrupt = Path(directory) / "corrupt.png"
            corrupt.write_bytes(b"not an image")
            for value in ("", Path(directory) / "missing.png", corrupt):
                with self.subTest(value=value):
                    presentation, _ = self._generate("{{image}}", value)
                    self.assertEqual(_picture_shapes(presentation), [])


class ImageLoopTest(unittest.TestCase):
    def test_image_and_table_loops_on_same_slide_are_rejected(self) -> None:
        with TemporaryDirectory() as directory:
            template = Path(directory) / "template.pptx"
            output = Path(directory) / "output.pptx"
            presentation = Presentation(TEMPLATE)
            slide = presentation.slides[0]
            table = slide.shapes.add_table(2, 1, 0, 0, 1000000, 1000000)
            table.name = "<foreach item=tables block=tableBlock></foreach>"
            presentation.save(template)
            with self.assertRaisesRegex(TemplateError, "cannot mix"):
                generate_presentation(
                    template,
                    {"imgs": [], "imageGroups": [], "tables": []},
                    output,
                )

    def test_empty_arrays_remove_loop_placeholders(self) -> None:
        with TemporaryDirectory() as directory:
            output = Path(directory) / "output.pptx"
            generate_presentation(
                TEMPLATE,
                {"imgs": [], "imageGroups": []},
                output,
            )
            presentation = Presentation(output)
            self.assertEqual(_picture_shapes(presentation), [])
            self.assertTrue(
                all(
                    "<foreach" not in shape.name
                    for slide in presentation.slides
                    for shape in slide.shapes
                )
            )

    def test_string_array_auto_grid_skips_invalid_images(self) -> None:
        with TemporaryDirectory() as directory:
            output = Path(directory) / "output.pptx"
            generate_presentation(
                TEMPLATE,
                {
                    "imgs": [
                        str(IMAGES / "wide.png"),
                        "",
                        str(IMAGES / "missing.png"),
                        str(IMAGES / "tall.png"),
                        str(IMAGES / "square.png"),
                    ],
                    "imageGroups": [],
                },
                output,
            )
            presentation = Presentation(output)
            pictures = _picture_shapes(presentation)
            self.assertEqual(len(pictures), 3)
            self.assertTrue(
                all(
                    picture.left + picture.width <= presentation.slide_width
                    for picture in pictures
                )
            )
            self.assertTrue(
                all(
                    picture.top + picture.height <= presentation.slide_height
                    for picture in pictures
                )
            )
            self.assertEqual(len({picture.name for picture in pictures}), 3)

    def test_explicit_rows_and_cols_control_page_capacity(self) -> None:
        with TemporaryDirectory() as directory:
            template = Path(directory) / "template.pptx"
            output = Path(directory) / "output.pptx"
            _loop_template(
                template,
                "<foreach item=imgs rows=1 cols=1></foreach>(x,y)",
            )
            generate_presentation(
                template,
                {
                    "imgs": [str(IMAGES / "square.png")] * 3,
                    "imageGroups": [],
                },
                output,
            )
            presentation = Presentation(output)
            image_pages = [
                slide
                for slide in presentation.slides
                if any(shape.shape_type == 13 for shape in slide.shapes)
            ]
            self.assertEqual(len(image_pages), 3)
            self.assertTrue(
                all(
                    sum(shape.shape_type == 13 for shape in slide.shapes) == 1
                    for slide in image_pages
                )
            )

    def test_rows_only_keeps_one_grid_row_and_auto_calculates_columns(self) -> None:
        with TemporaryDirectory() as directory:
            template = Path(directory) / "template.pptx"
            output = Path(directory) / "output.pptx"
            _loop_template(template, "<foreach item=imgs rows=1></foreach>(x,y)")
            generate_presentation(
                template,
                {
                    "imgs": [str(IMAGES / "square.png")] * 4,
                    "imageGroups": [],
                },
                output,
            )
            presentation = Presentation(output)
            for slide in presentation.slides:
                pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
                if pictures:
                    self.assertEqual(len({picture.top for picture in pictures}), 1)

    def test_cols_only_keeps_one_grid_column_and_auto_calculates_rows(self) -> None:
        with TemporaryDirectory() as directory:
            template = Path(directory) / "template.pptx"
            output = Path(directory) / "output.pptx"
            _loop_template(template, "<foreach item=imgs cols=1></foreach>(x,y)")
            generate_presentation(
                template,
                {
                    "imgs": [str(IMAGES / "square.png")] * 4,
                    "imageGroups": [],
                },
                output,
            )
            presentation = Presentation(output)
            for slide in presentation.slides:
                pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
                if pictures:
                    self.assertEqual(len({picture.left for picture in pictures}), 1)

    def test_group_loop_replaces_text_and_skips_invalid_whole_group(self) -> None:
        with TemporaryDirectory() as directory:
            output = Path(directory) / "output.pptx"
            generate_presentation(
                TEMPLATE,
                {
                    "imgs": [],
                    "imageGroups": [
                        {"title": "有效 A", "imagePath": str(IMAGES / "wide.png")},
                        {"title": "无效", "imagePath": str(IMAGES / "missing.png")},
                        {"title": "有效 B", "imagePath": str(IMAGES / "tall.png")},
                    ],
                },
                output,
            )
            presentation = Presentation(output)
            text = "\n".join(
                shape.text
                for slide in presentation.slides
                for shape in slide.shapes
                if getattr(shape, "has_text_frame", False)
            )
            self.assertIn("分类：有效 A", text)
            self.assertIn("分类：有效 B", text)
            self.assertNotIn("分类：无效", text)
            self.assertEqual(len(_picture_shapes(presentation)), 2)

    def test_image_limit_is_enforced(self) -> None:
        with TemporaryDirectory() as directory:
            with self.assertRaisesRegex(TemplateError, "image count exceeds"):
                generate_presentation(
                    TEMPLATE,
                    {
                        "imgs": [str(IMAGES / "square.png")] * 2,
                        "imageGroups": [],
                    },
                    Path(directory) / "output.pptx",
                    config=GeneratorConfig(max_images=1),
                )

    def test_total_image_bytes_limit_is_enforced(self) -> None:
        with TemporaryDirectory() as directory:
            with self.assertRaisesRegex(TemplateError, "total image bytes exceed"):
                generate_presentation(
                    TEMPLATE,
                    {
                        "imgs": [str(IMAGES / "square.png")],
                        "imageGroups": [],
                    },
                    Path(directory) / "output.pptx",
                    config=GeneratorConfig(max_total_image_bytes=1),
                )

    def test_allowed_image_roots_remove_out_of_scope_images(self) -> None:
        with TemporaryDirectory() as directory:
            output = Path(directory) / "output.pptx"
            generate_presentation(
                TEMPLATE,
                {
                    "imgs": [str(IMAGES / "square.png")],
                    "imageGroups": [],
                },
                output,
                config=GeneratorConfig(allowed_image_roots=(directory,)),
            )
            self.assertEqual(_picture_shapes(Presentation(output)), [])

    def test_image_file_and_pixel_limits_remove_oversized_images(self) -> None:
        with TemporaryDirectory() as directory:
            for config in (
                GeneratorConfig(max_image_bytes=1),
                GeneratorConfig(max_image_pixels=1),
            ):
                with self.subTest(config=config):
                    output = Path(directory) / f"output-{config.max_image_bytes}.pptx"
                    generate_presentation(
                        TEMPLATE,
                        {
                            "imgs": [str(IMAGES / "square.png")],
                            "imageGroups": [],
                        },
                        output,
                        config=config,
                    )
                    self.assertEqual(_picture_shapes(Presentation(output)), [])

    def test_rotated_loop_prototype_is_rejected_without_overwriting_output(self) -> None:
        with TemporaryDirectory() as directory:
            template = Path(directory) / "template.pptx"
            output = Path(directory) / "output.pptx"
            presentation = Presentation(TEMPLATE)
            picture = next(
                shape
                for shape in presentation.slides[0].shapes
                if shape.shape_type == 13
            )
            picture.rotation = 15
            presentation.save(template)
            output.write_bytes(b"existing")
            with self.assertRaisesRegex(TemplateError, "rotated prototypes"):
                generate_presentation(
                    template,
                    {"imgs": [str(IMAGES / "square.png")], "imageGroups": []},
                    output,
                )
            self.assertEqual(output.read_bytes(), b"existing")

    def test_single_image_replacement_can_coexist_with_table_generation(self) -> None:
        with TemporaryDirectory() as directory:
            template = Path(directory) / "template.pptx"
            output = Path(directory) / "output.pptx"
            presentation = Presentation(FIXTURES / "test2.pptx")
            logo = presentation.slides[0].shapes.add_picture(
                str(IMAGES / "square.png"), 0, 0, width=200_000, height=200_000
            )
            logo.name = "{{logo}}(x,y)"
            title = presentation.slides[0].shapes.add_textbox(
                300_000, 0, 2_000_000, 300_000
            )
            title.text = "报告：{{reportTitle}}"
            presentation.save(template)
            generate_presentation(
                template,
                {
                    "logo": str(IMAGES / "wide.png"),
                    "reportTitle": "统一入口",
                    "tables": [
                        {
                            "title": "混合内容",
                            "columns": ["A"],
                            "rows": [[1]],
                        }
                    ],
                },
                output,
            )
            generated = Presentation(output)
            self.assertTrue(
                any(
                    shape.has_table
                    for slide in generated.slides
                    for shape in slide.shapes
                )
            )
            self.assertTrue(
                any(
                    shape.name == "logo__image"
                    for slide in generated.slides
                    for shape in slide.shapes
                )
            )
            self.assertTrue(
                any(
                    getattr(shape, "has_text_frame", False)
                    and "报告：统一入口" in shape.text
                    for slide in generated.slides
                    for shape in slide.shapes
                )
            )

    def test_image_limit_counts_single_and_loop_images_together(self) -> None:
        with TemporaryDirectory() as directory:
            template = Path(directory) / "template.pptx"
            output = Path(directory) / "output.pptx"
            _single_image_template(template, "{{logo}}(x,y)")
            with self.assertRaisesRegex(TemplateError, "image count exceeds"):
                generate_presentation(
                    template,
                    {
                        "logo": str(IMAGES / "square.png"),
                        "imageGroups": [
                            {
                                "title": "设备",
                                "imagePath": str(IMAGES / "square.png"),
                            }
                        ],
                    },
                    output,
                    config=GeneratorConfig(max_images=1),
                )

    def test_malformed_single_image_directive_is_rejected(self) -> None:
        with TemporaryDirectory() as directory:
            template = Path(directory) / "template.pptx"
            output = Path(directory) / "output.pptx"
            _single_image_template(template, "{{logo}}(invalid)")
            with self.assertRaisesRegex(TemplateError, "malformed image"):
                generate_presentation(
                    template,
                    {"logo": str(IMAGES / "square.png"), "imageGroups": []},
                    output,
                )


if __name__ == "__main__":
    unittest.main()
