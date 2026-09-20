from __future__ import annotations

import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

from genppt import TemplateError, generate_presentation


def _blank_presentation() -> tuple[Presentation, object]:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    return presentation, slide


class TextGenerationTest(unittest.TestCase):
    def test_plain_and_nested_text_expressions_are_replaced(self) -> None:
        with TemporaryDirectory() as directory:
            template = Path(directory) / "template.pptx"
            output = Path(directory) / "output.pptx"
            presentation, slide = _blank_presentation()
            box = slide.shapes.add_textbox(0, 0, Inches(6), Inches(1))
            box.text = "报告：{{report.title}}；缺失：{{missing}}"
            presentation.save(template)
            generate_presentation(
                template,
                {"report": {"title": "经营分析"}},
                output,
            )
            generated = Presentation(output)
            text = next(shape.text for shape in generated.slides[0].shapes)
            self.assertEqual(text, "报告：经营分析；缺失：")

    def test_expression_split_across_runs_preserves_surrounding_styles(self) -> None:
        with TemporaryDirectory() as directory:
            template = Path(directory) / "template.pptx"
            output = Path(directory) / "output.pptx"
            presentation, slide = _blank_presentation()
            box = slide.shapes.add_textbox(0, 0, Inches(6), Inches(1))
            paragraph = box.text_frame.paragraphs[0]
            paragraph.clear()
            first = paragraph.add_run()
            first.text = "前缀 {{na"
            first.font.bold = True
            first.font.color.rgb = RGBColor(255, 0, 0)
            second = paragraph.add_run()
            second.text = "me}} 后缀"
            second.font.italic = True
            presentation.save(template)
            generate_presentation(template, {"name": "张三"}, output)
            generated = Presentation(output)
            result = generated.slides[0].shapes[0].text_frame.paragraphs[0]
            self.assertEqual("".join(run.text for run in result.runs), "前缀 张三 后缀")
            self.assertTrue(result.runs[0].font.bold)
            self.assertEqual(result.runs[0].font.color.rgb, RGBColor(255, 0, 0))
            self.assertTrue(result.runs[1].font.italic)

    def test_static_table_cells_are_replaced(self) -> None:
        with TemporaryDirectory() as directory:
            template = Path(directory) / "template.pptx"
            output = Path(directory) / "output.pptx"
            presentation, slide = _blank_presentation()
            table = slide.shapes.add_table(1, 2, 0, 0, Inches(5), Inches(1)).table
            table.cell(0, 0).text = "{{left}}"
            table.cell(0, 1).text = "{{right}}"
            presentation.save(template)
            generate_presentation(template, {"left": "甲", "right": "乙"}, output)
            generated_table = Presentation(output).slides[0].shapes[0].table
            self.assertEqual(generated_table.cell(0, 0).text, "甲")
            self.assertEqual(generated_table.cell(0, 1).text, "乙")

    def test_template_without_directives_is_copied_successfully(self) -> None:
        with TemporaryDirectory() as directory:
            template = Path(directory) / "template.pptx"
            output = Path(directory) / "output.pptx"
            presentation, slide = _blank_presentation()
            slide.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = "固定内容"
            presentation.save(template)
            generate_presentation(template, {}, output)
            self.assertEqual(Presentation(output).slides[0].shapes[0].text, "固定内容")

    def test_unclaimed_text_foreach_is_rejected_instead_of_silently_ignored(self) -> None:
        with TemporaryDirectory() as directory:
            template = Path(directory) / "template.pptx"
            output = Path(directory) / "output.pptx"
            presentation, slide = _blank_presentation()
            shape = slide.shapes.add_textbox(0, 0, Inches(4), Inches(1))
            shape.name = "<foreach item=items>{{name}}</foreach>"
            shape.text = "{{name}}"
            presentation.save(template)
            with self.assertRaisesRegex(TemplateError, "unsupported foreach"):
                generate_presentation(template, {"items": []}, output)

    def test_malformed_foreach_is_rejected_during_template_planning(self) -> None:
        with TemporaryDirectory() as directory:
            template = Path(directory) / "template.pptx"
            output = Path(directory) / "output.pptx"
            presentation, slide = _blank_presentation()
            shape = slide.shapes.add_textbox(0, 0, Inches(4), Inches(1))
            shape.name = "<foreach item=items>{{name}}"
            presentation.save(template)
            with self.assertRaisesRegex(TemplateError, "unsupported foreach"):
                generate_presentation(template, {"items": []}, output)


if __name__ == "__main__":
    unittest.main()
