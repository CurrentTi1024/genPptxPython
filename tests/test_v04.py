from __future__ import annotations

import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from genppt import GeneratorConfig, TemplateError, generate_presentation
from genppt.data import normalize_table
from genppt.directives import parse_foreach, parse_single_image

IMAGE = Path(__file__).parent / "fixtures/images/square.png"


def deck():
    p = Presentation()
    p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
    return p


def text(slide, value, top=1, name="Text", left=0.5, width=5, height=0.4):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    shape.text = value
    shape.name = name
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(16)
    return shape


def table(slide, item="items", block="g", field="table1", top=1.6, left=0.5):
    shape = slide.shapes.add_table(2, 2, Inches(left), Inches(top), Inches(5), Inches(0.8))
    shape.name = f"<foreach item={item} block={block}>{{{{{field}}}}}</foreach>"
    shape.table.cell(0, 0).text = "Prototype"
    shape.table.cell(0, 0).text_frame.paragraphs[0].runs[0].font.bold = True
    return shape


def picture(slide, field="images", top=2.6, left=0.5, extra=""):
    shape = slide.shapes.add_picture(
        str(IMAGE), Inches(left), Inches(top), width=Inches(1), height=Inches(1)
    )
    shape.name = f"<foreach item=items block=g {extra}>{{{{{field}}}}}</foreach>（x，y）"
    return shape


def all_text(p):
    return [s.text for slide in p.slides for s in slide.shapes if s.has_text_frame]


def all_tables(p):
    return [s for slide in p.slides for s in slide.shapes if s.has_table]


class V04Test(unittest.TestCase):
    def test_inherited_style_merge_obeys_exclusive_choices_and_schema_order(self):
        from genppt.pptx_utils import _merge_properties

        base, overlay = OxmlElement("a:rPr"), OxmlElement("a:rPr")
        base.append(OxmlElement("a:latin"))
        overlay.append(OxmlElement("a:solidFill"))
        result = _merge_properties(base, overlay)
        self.assertEqual([n.tag.rsplit("}", 1)[-1] for n in result], ["solidFill", "latin"])
        for kind, first, second in (
            ("bodyPr", "normAutofit", "spAutoFit"),
            ("pPr", "buClr", "buClrTx"),
            ("pPr", "buSzPts", "buSzTx"),
        ):
            base, overlay = OxmlElement(f"a:{kind}"), OxmlElement(f"a:{kind}")
            base.append(OxmlElement(f"a:{first}"))
            overlay.append(OxmlElement(f"a:{second}"))
            self.assertEqual(len(_merge_properties(base, overlay)), 1)

    def test_empty_rows_delete_table_even_when_columns_are_empty(self):
        p = deck()
        slide = p.slides.add_slide(p.slide_layouts[6])
        table(slide)
        result = self.generate(p, {"items": [{"table1": {"columns": [], "rows": []}}]})
        self.assertEqual(all_tables(result), [])

    def test_invalid_table_selector_and_fractional_repeat_columns_are_rejected(self):
        p = deck()
        slide = p.slides.add_slide(p.slide_layouts[6])
        shape = table(slide)
        shape.name = "<foreach item=items>{{table1}}suffix</foreach>"
        with self.assertRaisesRegex(TemplateError, "table selector"):
            self.generate(p, {})
        for value in (1.5, True, "2", float("inf")):
            with self.subTest(value=value), self.assertRaises(TemplateError):
                normalize_table(
                    {"rows": [{"A": 1}], "_table": {"repeatLeadingColumns": value}},
                    GeneratorConfig(),
                )

    def test_static_field_nodes_are_preserved_verbatim(self):
        p = deck()
        slide = p.slides.add_slide(p.slide_layouts[6])
        shape = text(slide, "")
        paragraph = shape.text_frame.paragraphs[0]
        field = OxmlElement("a:fld")
        field.set("id", "{C5E94150-5D98-4A37-91EE-880BD5BDECE1}")
        field.set("type", "slidenum")
        value = OxmlElement("a:t")
        value.text = "1"
        field.append(value)
        paragraph._p.append(field)
        source = paragraph._p.xml
        result = self.generate(p, {})
        self.assertEqual(result.slides[0].shapes[0].text_frame.paragraphs[0]._p.xml, source)

    def test_template_package_limits_and_corrupt_input_preserve_output(self):
        p = deck()
        p.slides.add_slide(p.slide_layouts[6])
        with TemporaryDirectory() as directory:
            template, output = Path(directory) / "in.pptx", Path(directory) / "out.pptx"
            p.save(template)
            original = template.read_bytes()
            output.write_bytes(b"keep")
            for config in (
                {"max_template_bytes": 1},
                {"max_template_parts": 1},
                {"max_template_uncompressed_bytes": 1},
            ):
                with self.subTest(config=config), self.assertRaises(TemplateError):
                    generate_presentation(template, {}, output, config=GeneratorConfig(**config))
                self.assertEqual(output.read_bytes(), b"keep")
                self.assertEqual(template.read_bytes(), original)
            template.write_bytes(b"not a zip")
            with self.assertRaisesRegex(TemplateError, "valid PowerPoint package"):
                generate_presentation(template, {}, output)
            self.assertEqual(output.read_bytes(), b"keep")

    def test_repeated_header_output_cells_are_bounded(self):
        p = deck()
        slide = p.slides.add_slide(p.slide_layouts[6])
        table(slide)
        with self.assertRaisesRegex(TemplateError, "generated cell count"):
            self.generate(
                p,
                {"items": [{"table1": {"rows": [{"A": i} for i in range(40)]}}]},
                max_total_cells=41,
            )

    def test_long_block_names_cannot_collide_after_truncation(self):
        p = deck()
        slide = p.slides.add_slide(p.slide_layouts[6])
        for index in range(2):
            text(
                slide,
                "{{title}}",
                top=1 + index * 2,
                name=f"<foreach item=items block={'z' * 45}{index}>Caption</foreach>",
            )
        result = self.generate(p, {"items": [{"title": "A"}]})
        names = [s.name for s in result.slides[0].shapes]
        self.assertEqual(len(names), len(set(names)))

    def test_save_failure_preserves_output_and_cleans_temporary_package(self):
        from unittest.mock import patch

        p = deck()
        p.slides.add_slide(p.slide_layouts[6])
        with TemporaryDirectory() as directory:
            template, output = Path(directory) / "in.pptx", Path(directory) / "out.pptx"
            p.save(template)
            output.write_bytes(b"keep")
            with (
                patch("pptx.presentation.Presentation.save", side_effect=OSError("disk full")),
                self.assertRaises(OSError),
            ):
                generate_presentation(template, {}, output)
            self.assertEqual(output.read_bytes(), b"keep")
            self.assertEqual(list(Path(directory).glob(".*.tmp.pptx")), [])

    def test_parallel_requests_have_isolated_context_and_image_caches(self):
        from concurrent.futures import ThreadPoolExecutor

        p = deck()
        slide = p.slides.add_slide(p.slide_layouts[6])
        text(slide, "{{title}}", name="<foreach item=items block=g>Caption</foreach>")
        table(slide)
        picture(slide)
        with TemporaryDirectory() as directory:
            template = Path(directory) / "in.pptx"
            p.save(template)

            def run(index):
                output = Path(directory) / f"out{index}.pptx"
                generate_presentation(
                    template,
                    {"items": [{"title": str(index), "images": [str(IMAGE)]}]},
                    output,
                )
                return all_text(Presentation(output))

            with ThreadPoolExecutor(max_workers=4) as pool:
                results = list(pool.map(run, range(8)))
            self.assertEqual(results, [[str(i)] for i in range(8)])

    def test_minimum_fragment_rows_moves_tiny_initial_fragment_to_fresh_page(self):
        p = deck()
        slide = p.slides.add_slide(p.slide_layouts[6])
        table(slide, top=6.05)
        result = self.generate(p, {"items": [{"table1": {"rows": [{"A": i} for i in range(40)]}}]})
        self.assertFalse(any(s.has_table for s in result.slides[0].shapes))
        self.assertEqual(sum(len(s.table.rows) - 1 for s in all_tables(result)), 40)

    def test_real_template_end_to_end_data_and_page_order(self):
        from examples.run_multipage_demo import demo_payload

        p = Presentation(Path(__file__).parent / "fixtures/test-multipage.pptx")
        result = self.generate(p, demo_payload())
        self.assertEqual(len(result.slides), 6)
        native = all_tables(result)
        self.assertEqual(len(native), 7)
        self.assertEqual(sum(len(s.table.rows) - 1 for s in native), 31)
        self.assertEqual(
            sum(s.shape_type == 13 for slide in result.slides for s in slide.shapes), 4
        )
        self.assertEqual(result.slides[0].shapes[0].text, "CPS")
        self.assertEqual(result.slides[1].shapes[0].text, "CNS")
        self.assertEqual(result.slides[-1].shapes[0].text, "Cns overall")
        rows = [
            s.table.cell(r, 0).text
            for s in native
            if s.table.cell(0, 0).text == "ID"
            for r in range(1, len(s.table.rows))
        ]
        self.assertEqual(rows, [str(i) for i in range(1, 26)])

    def test_effective_placeholder_style_is_identical_on_blank_continuation(self):
        p = Presentation(Path(__file__).parent / "fixtures/test-multipage.pptx")
        result = self.generate(p, {"tables": [{"title": "A"}, {"title": "B"}, {"title": "C"}]})
        # Force separate pages without altering the shared source template.
        caption = next(s for s in p.slides[1].shapes if s.has_text_frame and "foreach" in s.name)
        caption.name = "<foreach item=tables block=reasonTable page=group>Caption</foreach>"
        result = self.generate(p, {"tables": [{"title": "A"}, {"title": "B"}]})
        captions = [
            s
            for slide in result.slides
            for s in slide.shapes
            if s.has_text_frame and "分类：" in s.text
        ]
        self.assertEqual(len(captions), 2)
        for shape in captions:
            self.assertFalse(shape.is_placeholder)
            paragraph = shape.text_frame.paragraphs[0]
            self.assertEqual(paragraph.runs[0].font.size, Pt(28))
            self.assertTrue(paragraph.runs[0].font.bold)
            self.assertEqual(paragraph._p.xpath("./a:pPr/a:buChar")[0].get("char"), "l")
        self.assertEqual(
            captions[0].text_frame.paragraphs[0]._p.pPr.xml,
            captions[1].text_frame.paragraphs[0]._p.pPr.xml,
        )

    def test_image_exif_orientation_and_webp_are_normalized(self):
        with TemporaryDirectory() as directory:
            path = Path(directory)
            for filename in ("oriented.jpg", "image.webp"):
                image = Image.new("RGB", (20, 40), "red")
                if filename.endswith("jpg"):
                    exif = image.getexif()
                    exif[274] = 6
                    image.save(path / filename, exif=exif)
                else:
                    image.save(path / filename)
                p = deck()
                slide = p.slides.add_slide(p.slide_layouts[6])
                shape = picture(slide)
                shape.name = "{{photo}}(x)"
                result = self.generate(p, {"photo": str(path / filename)})
                photo = result.slides[0].shapes[0]
                expected = 2 if filename.endswith("jpg") else 0.5
                self.assertAlmostEqual(photo.width / photo.height, expected)

    def test_nonfinite_config_decimal_and_excessive_format_are_rejected(self):
        from decimal import Decimal

        for config in (
            {"block_gap_points": float("nan")},
            {"max_tables": float("inf")},
            {"max_images": True},
        ):
            with self.assertRaises(ValueError):
                GeneratorConfig(**config)
        for value, formats in ((Decimal("NaN"), {}), (1, {"A": "1000000000d"})):
            with self.assertRaises(TemplateError):
                normalize_table(
                    {"rows": [{"A": value}], "_table": {"formats": formats}}, GeneratorConfig()
                )

    def test_explicit_nested_table_selector_precedes_legacy_rows(self):
        p = deck()
        slide = p.slides.add_slide(p.slide_layouts[6])
        table(slide)
        result = self.generate(
            p, {"items": [{"rows": [{"A": "outer"}], "table1": {"rows": [{"A": "inner"}]}}]}
        )
        self.assertEqual(all_tables(result)[0].table.cell(1, 0).text, "inner")

    def test_resource_failure_preserves_existing_output_and_removes_temp_file(self):
        p = deck()
        for _ in range(2):
            slide = p.slides.add_slide(p.slide_layouts[6])
            table(slide)
        with TemporaryDirectory() as directory:
            template, output = Path(directory) / "in.pptx", Path(directory) / "out.pptx"
            p.save(template)
            output.write_bytes(b"original")
            with self.assertRaises(TemplateError):
                generate_presentation(
                    template,
                    {"items": [{"table1": {"rows": [{"A": 1}]}}]},
                    output,
                    config=GeneratorConfig(max_tables=1),
                )
            self.assertEqual(output.read_bytes(), b"original")
            self.assertEqual(list(Path(directory).glob("*.tmp.pptx")), [])

    def generate(self, presentation, data, **config):
        directory = TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        template, output = Path(directory.name) / "in.pptx", Path(directory.name) / "out.pptx"
        presentation.save(template)
        generate_presentation(template, data, output, config=GeneratorConfig(**config))
        return Presentation(output)

    def test_all_original_pages_expand_in_order_with_static_page(self):
        p = deck()
        for index in range(3):
            slide = p.slides.add_slide(p.slide_layouts[6])
            text(slide, f"Static {index}", top=0.1)
            if index != 1:
                shape = table(slide, item=f"items{index}", field="", top=1)
                shape.name = f"<foreach item=items{index} page=group></foreach>"
        result = self.generate(
            p,
            {
                "items0": [{"rows": [{"A": "first"}]}, {"rows": [{"A": "second"}]}],
                "items2": [{"rows": [{"A": "third"}]}],
            },
        )
        self.assertEqual(len(result.slides), 4)
        self.assertIn("Static 0", all_text(result))
        self.assertEqual(len(result.slides[1].shapes), 1)
        self.assertEqual(result.slides[2].shapes[0].text, "Static 1")
        self.assertEqual(
            [t.table.cell(1, 0).text for t in all_tables(result)], ["first", "second", "third"]
        )

    def test_same_block_name_on_two_pages_is_scoped_to_page(self):
        p = deck()
        for index in range(2):
            slide = p.slides.add_slide(p.slide_layouts[6])
            table(slide, left=0.5 + index)
        result = self.generate(p, {"items": [{"table1": {"rows": [{"A": 1}]}}]})
        self.assertEqual(len(result.slides), 2)
        self.assertEqual([t.left for t in all_tables(result)], [Inches(0.5), Inches(1.5)])

    def test_missing_outer_array_keeps_static_original_page(self):
        p = deck()
        slide = p.slides.add_slide(p.slide_layouts[6])
        text(slide, "CPS", top=0.1)
        table(slide)
        result = self.generate(p, {})
        self.assertEqual(all_text(result), ["CPS"])
        self.assertEqual(len(result.slides), 1)

    def test_mixed_group_long_table_finishes_before_pictures(self):
        p = deck()
        slide = p.slides.add_slide(p.slide_layouts[6])
        text(slide, "Original title", top=0.1)
        text(slide, "Group {{title}}", name="<foreach item=items block=g>Title</foreach>")
        table(slide)
        picture(slide, field="a")
        picture(slide, field="b", left=3, top=2.65)
        result = self.generate(
            p,
            {
                "items": [
                    {
                        "title": "A",
                        "table1": {
                            "columns": [{"n": "Number"}],
                            "rows": [{"n": i} for i in range(40)],
                        },
                        "a": str(IMAGE),
                        "b": str(IMAGE),
                    }
                ]
            },
        )
        self.assertGreater(len(result.slides), 2)
        self.assertEqual(sum("Group A" in v for v in all_text(result)), 1)
        self.assertEqual(sum("Original title" in v for v in all_text(result)), 1)
        self.assertEqual(sum(len(t.table.rows) - 1 for t in all_tables(result)), 40)
        self.assertTrue(all(t.table.cell(0, 0).text == "Number" for t in all_tables(result)))
        table_pages = [
            i for i, slide in enumerate(result.slides) if any(s.has_table for s in slide.shapes)
        ]
        picture_pages = [
            i
            for i, slide in enumerate(result.slides)
            if any(s.shape_type == 13 for s in slide.shapes)
        ]
        self.assertGreaterEqual(min(picture_pages), max(table_pages))
        last = result.slides[picture_pages[-1]]
        photos = [s for s in last.shapes if s.shape_type == 13]
        self.assertEqual(len(photos), 2)
        self.assertEqual(photos[1].top - photos[0].top, Inches(0.05))
        self.assertIs(
            result.slides[1].slide_layout.slide_master, result.slides[0].slide_layout.slide_master
        )

    def test_multiple_table_fields_and_array_in_one_group(self):
        p = deck()
        slide = p.slides.add_slide(p.slide_layouts[6])
        table(slide)
        table(slide, field="table2", top=2.8)
        result = self.generate(
            p,
            {
                "items": [
                    {
                        "table1": {"rows": [{"A": 1}]},
                        "table2": [{"rows": [{"B": 2}]}, {"rows": [{"C": 3}]}],
                    }
                ]
            },
        )
        self.assertEqual([t.table.cell(1, 0).text for t in all_tables(result)], ["1", "2", "3"])

    def test_empty_member_keeps_title_and_other_picture_and_compacts(self):
        p = deck()
        slide = p.slides.add_slide(p.slide_layouts[6])
        text(slide, "{{title}}", name="<foreach item=items block=g>Title</foreach>")
        table(slide)
        picture(slide, "a")
        picture(slide, "b", left=3)
        result = self.generate(p, {"items": [{"title": "Keep", "b": str(IMAGE)}]})
        self.assertEqual(all_text(result), ["Keep"])
        self.assertEqual(len(all_tables(result)), 0)
        photos = [s for s in result.slides[0].shapes if s.shape_type == 13]
        self.assertEqual(len(photos), 1)
        self.assertLess(photos[0].top, Inches(2))

    def test_group_page_mode_and_literal_text_shape_name(self):
        p = deck()
        slide = p.slides.add_slide(p.slide_layouts[6])
        text(slide, "{{title}}", name="<foreach item=items block=g page=group>{{name}}</foreach>")
        result = self.generate(p, {"items": [{"title": "A", "name": "Not a name"}, {"title": "B"}]})
        self.assertEqual(len(result.slides), 2)
        self.assertTrue(result.slides[0].shapes[0].name.startswith("{{name}}__"))
        self.assertEqual(result.slides[1].shapes[0].top, Inches(0.55))

    def test_standalone_text_shape_loop_is_vertical(self):
        p = deck()
        slide = p.slides.add_slide(p.slide_layouts[6])
        text(slide, "{{title}}", name="<foreach item=items>Label</foreach>")
        result = self.generate(p, {"items": [{"title": "A"}, {"title": "B"}]})
        self.assertEqual(all_text(result), ["A", "B"])
        self.assertGreater(result.slides[0].shapes[1].top, result.slides[0].shapes[0].top)

    def test_inline_if_false_values_and_missing_placeholder(self):
        for value in (None, False, 0, "", [], {}):
            with self.subTest(value=value):
                p = deck()
                slide = p.slides.add_slide(p.slide_layouts[6])
                text(slide, "A{{missing}}<if status>B{{status}}</if>C")
                self.assertEqual(all_text(self.generate(p, {"status": value})), ["AC"])

    def test_paragraph_loop_preserves_bullets_spacing_and_first_brace_style(self):
        p = deck()
        slide = p.slides.add_slide(p.slide_layouts[6])
        shape = text(slide, "", height=0.4)
        paragraph = shape.text_frame.paragraphs[0]
        paragraph.level = 1
        paragraph.space_after = Pt(9)
        bullet = OxmlElement("a:buChar")
        bullet.set("char", "•")
        paragraph._p.get_or_add_pPr().append(bullet)
        paragraph.clear()
        a = paragraph.add_run()
        a.text = "<foreach items=details>Value: {"
        a.font.bold = True
        a.font.size = Pt(18)
        a.font.color.rgb = RGBColor(200, 0, 0)
        b = paragraph.add_run()
        b.text = "{name}}<if status> OK</if></foreach>"
        b.font.italic = True
        result = self.generate(p, {"details": [{"name": "X", "status": 1}, {"name": "Y"}]})
        paragraphs = result.slides[0].shapes[0].text_frame.paragraphs
        self.assertEqual([v.text for v in paragraphs], ["Value: X OK", "Value: Y"])
        for para, value in zip(paragraphs, ("X", "Y"), strict=True):
            self.assertEqual(para.level, 1)
            self.assertEqual(para.space_after, Pt(9))
            self.assertEqual(para._p.xpath("./a:pPr/a:buChar")[0].get("char"), "•")
            run = next(r for r in para.runs if r.text == value)
            self.assertTrue(run.font.bold)
            self.assertEqual(run.font.size, Pt(18))
            self.assertEqual(run.font.color.rgb, RGBColor(200, 0, 0))

    def test_empty_paragraph_loop_and_conditional_shape_are_removed(self):
        p = deck()
        slide = p.slides.add_slide(p.slide_layouts[6])
        text(slide, "<foreach items=details>{{name}}</foreach>")
        text(slide, "<if status>{{status}}</if>", top=2)
        result = self.generate(p, {})
        self.assertEqual(len(result.slides[0].shapes), 0)

    def test_text_value_is_never_reinterpreted_as_tags(self):
        p = deck()
        slide = p.slides.add_slide(p.slide_layouts[6])
        text(slide, "{{title}}")
        value = "{{other}}<if nope>Literal</if>"
        self.assertEqual(all_text(self.generate(p, {"title": value})), [value])

    def test_text_foreach_requires_object_array_and_whole_paragraph(self):
        for template, data in [
            ("<foreach items=details>{{name}}</foreach>", {"details": [1]}),
            ("prefix <foreach items=details>{{name}}</foreach>", {}),
            ("<foreach items=details><foreach items=nested>{{name}}</foreach></foreach>", {}),
            ("<if status>unclosed", {}),
        ]:
            with self.subTest(template=template):
                p = deck()
                slide = p.slides.add_slide(p.slide_layouts[6])
                text(slide, template)
                with self.assertRaises(TemplateError):
                    self.generate(p, data)

    def test_text_grid_attributes_are_rejected(self):
        p = deck()
        slide = p.slides.add_slide(p.slide_layouts[6])
        text(slide, "{{title}}", name="<foreach item=items rows=2>Title</foreach>")
        with self.assertRaisesRegex(TemplateError, "only supported on pictures"):
            self.generate(p, {})

    def test_chinese_brackets_and_commas(self):
        for suffix in ("（x，y）", "(x，y)", "（x,y）", "（y）"):
            self.assertIsNotNone(parse_single_image("{{image}}" + suffix))
            self.assertIsNotNone(parse_foreach("<foreach item=images></foreach>" + suffix, 1))

    def test_columns_are_inferred_only_from_first_row(self):
        normalized = normalize_table({"rows": [{"A": 1}, {"A": 2, "B": 3}]}, GeneratorConfig())
        self.assertEqual(normalized.labels, ["A"])
        self.assertEqual(normalized.rows, [["1"], ["2"]])

    def test_resource_counts_accumulate_across_original_pages(self):
        p = deck()
        for _ in range(2):
            slide = p.slides.add_slide(p.slide_layouts[6])
            table(slide)
        for config in ({"max_tables": 1}, {"max_total_cells": 3}, {"max_loop_items": 1}):
            with self.subTest(config=config), self.assertRaises(TemplateError):
                self.generate(p, {"items": [{"table1": {"rows": [{"A": 1}]}}]}, **config)

    def test_text_resource_caps(self):
        p = deck()
        slide = p.slides.add_slide(p.slide_layouts[6])
        text(slide, "<foreach items=details>{{name}}</foreach>")
        for config in (
            {"max_text_paragraphs": 1},
            {"max_text_characters": 2},
            {"max_loop_items": 1},
        ):
            with self.subTest(config=config), self.assertRaises(TemplateError):
                self.generate(p, {"details": [{"name": "AA"}, {"name": "BB"}]}, **config)

    def test_single_table_object_and_single_picture_path(self):
        p = deck()
        slide = p.slides.add_slide(p.slide_layouts[6])
        shape = table(slide, item="report", field="", top=1)
        shape.name = "<foreach item=report></foreach>"
        photo = picture(slide, top=3)
        photo.name = "<foreach item=photo></foreach>（x，y）"
        result = self.generate(p, {"report": {"rows": [{"A": 1}]}, "photo": str(IMAGE)})
        self.assertEqual(len(all_tables(result)), 1)
        self.assertEqual(sum(s.shape_type == 13 for s in result.slides[0].shapes), 1)

    def test_long_paragraph_list_continues_on_blank_master_pages(self):
        p = deck()
        slide = p.slides.add_slide(p.slide_layouts[6])
        text(slide, "Root title", top=0.1)
        text(slide, "<foreach items=details>{{name}}</foreach>")
        result = self.generate(p, {"details": [{"name": f"Line {i}"} for i in range(40)]})
        self.assertGreater(len(result.slides), 1)
        self.assertEqual(sum(v == "Root title" for v in all_text(result)), 1)
        joined = "\n".join(all_text(result))
        for i in range(40):
            self.assertIn(f"Line {i}", joined)

    def test_generated_names_are_unique_and_geometry_is_in_bounds(self):
        p = deck()
        for _ in range(2):
            slide = p.slides.add_slide(p.slide_layouts[6])
            table(slide)
            picture(slide, extra="rows=1 cols=2")
        result = self.generate(
            p,
            {
                "items": [
                    {"table1": {"rows": [{"A": i} for i in range(20)]}, "images": [str(IMAGE)] * 5}
                ]
            },
        )
        names = [s.name for slide in result.slides for s in slide.shapes]
        self.assertEqual(len(names), len(set(names)))
        for slide in result.slides:
            for shape in slide.shapes:
                self.assertGreaterEqual(shape.top, 0)
                self.assertLessEqual(shape.top + shape.height, result.slide_height)
                self.assertLessEqual(shape.left + shape.width, result.slide_width)


if __name__ == "__main__":
    unittest.main()
