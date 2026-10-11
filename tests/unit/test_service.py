from __future__ import annotations

import inspect
import json
import unittest
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO, StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZipFile

from pptx import Presentation

from genppt import GeneratorConfig, PresentationService, TemplateError, generate_ppt
from genppt.core.settings import Settings
from genppt.engine.package import read_template
from genppt.services.presentation import parse_data

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures"
TEMPLATE = FIXTURES / "test-multipage.pptx"


class ServiceTest(unittest.TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.settings = Settings(output_dir=Path(self.directory.name) / "output")
        self.service = PresentationService(self.settings)

    def test_exactly_two_argument_entry_returns_existing_absolute_path(self):
        self.assertEqual(list(inspect.signature(generate_ppt).parameters), ["template", "data"])
        self.assertEqual(
            list(inspect.signature(self.service.generate).parameters), ["template", "data"]
        )
        result = self.service.generate(str(TEMPLATE), "{}")
        self.assertIsInstance(result, str)
        self.assertTrue(Path(result).is_absolute())
        self.assertEqual(len(Presentation(result).slides), 3)

    def test_seekable_stream_is_rewound_restored_and_not_closed(self):
        stream = BytesIO(TEMPLATE.read_bytes())
        stream.seek(17)
        result = self.service.generate(stream, "{}")
        self.assertEqual(stream.tell(), 17)
        self.assertFalse(stream.closed)
        self.assertEqual(len(Presentation(result).slides), 3)

    def test_stream_position_is_restored_on_package_failure(self):
        stream = BytesIO(b"not a zip")
        stream.seek(2)
        with self.assertRaises(TemplateError):
            self.service.generate(stream, "{}")
        self.assertEqual(stream.tell(), 2)
        self.assertFalse(self.settings.output_dir.exists())

    def test_nonseekable_short_reads_are_supported(self):
        class Stream:
            def __init__(self):
                self.source = BytesIO(TEMPLATE.read_bytes())

            def read(self, count):
                return self.source.read(min(count, 127))

        self.assertEqual(len(Presentation(self.service.generate(Stream(), "{}")).slides), 3)

    def test_byte_buffer_is_supported(self):
        self.assertTrue(Path(self.service.generate(TEMPLATE.read_bytes(), "{}")).is_file())

    def test_invalid_stream_types_and_closed_stream_are_rejected(self):
        stream = BytesIO()
        stream.close()
        for value in (None, 123, StringIO("text"), stream):
            with self.subTest(value=value), self.assertRaises(TemplateError):
                self.service.generate(value, "{}")

    def test_stream_size_cap_and_pointer_restore(self):
        stream = BytesIO(b"abcdefgh")
        stream.seek(3)
        with self.assertRaises(TemplateError):
            read_template(stream, GeneratorConfig(max_template_bytes=4))
        self.assertEqual(stream.tell(), 3)

    def test_duplicate_package_parts_are_rejected(self):
        stream = BytesIO()
        with ZipFile(stream, "w") as archive:
            archive.writestr("duplicate", "A")
            with self.assertWarns(UserWarning):
                archive.writestr("duplicate", "B")
        with self.assertRaisesRegex(TemplateError, "duplicate template"):
            self.service.generate(stream, "{}")

    def test_invalid_json_roots_and_non_string_data_are_rejected(self):
        for value in (None, {}, "[]", "null", "1", '"string"', "{invalid}"):
            with self.subTest(value=value), self.assertRaises(TemplateError):
                self.service.generate(TEMPLATE, value)
        self.assertFalse(self.settings.output_dir.exists())

    def test_duplicate_json_keys_and_nonfinite_values_are_rejected(self):
        for value in ('{"a":1,"a":2}', '{"a":NaN}', '{"a":Infinity}', '{"a":1e999}'):
            with self.subTest(value=value), self.assertRaises(TemplateError):
                parse_data(value, self.settings)

    def test_xml_control_characters_and_surrogates_are_rejected(self):
        for value in ('{"a":"\\u0000"}', '{"a":"\\ud800"}', '{"a":"\\uffff"}'):
            with self.subTest(value=value), self.assertRaises(TemplateError):
                parse_data(value, self.settings)

    def test_json_size_depth_and_node_caps(self):
        for value, limits in (
            ('{"a":"汉字"}', {"max_data_bytes": 10}),
            ('{"a":{"b":1}}', {"max_json_depth": 2}),
            ('{"a":[1,2,3]}', {"max_json_nodes": 3}),
            ('{"a":' + "[" * 2000 + "0" + "]" * 2000 + "}", {}),
        ):
            with self.subTest(limits=limits), self.assertRaises(TemplateError):
                parse_data(value, Settings(**limits))

    def test_unicode_json_and_empty_object_work(self):
        value = {"title": "中文标题", "optional": None, "rows": []}
        self.assertEqual(parse_data(json.dumps(value), self.settings), value)
        self.assertEqual(parse_data("{}", self.settings), {})

    def test_parallel_calls_create_unique_paths(self):
        with ThreadPoolExecutor(max_workers=4) as pool:
            paths = list(pool.map(lambda _: self.service.generate(TEMPLATE, "{}"), range(8)))
        self.assertEqual(len(set(paths)), 8)
        self.assertTrue(all(Path(path).is_file() for path in paths))

    def test_failed_save_leaves_no_generated_or_temporary_file(self):
        with (
            patch("pptx.presentation.Presentation.save", side_effect=OSError("disk full")),
            self.assertRaises(OSError),
        ):
            self.service.generate(TEMPLATE, "{}")
        self.assertEqual(list(self.settings.output_dir.iterdir()), [])

    def test_environment_convenience_entry_and_unique_output(self):
        with patch.dict("os.environ", {"GENPPT_OUTPUT_DIR": self.directory.name}):
            result = generate_ppt(TEMPLATE, "{}")
        self.assertEqual(Path(result).parent, Path(self.directory.name).resolve())

    def test_invalid_settings_fail_fast(self):
        for values in (
            {"max_data_bytes": 0},
            {"max_json_nodes": True},
            {"max_json_depth": -1},
            {"max_concurrent_jobs": 1.5},
        ):
            with self.subTest(values=values), self.assertRaises(ValueError):
                Settings(**values)

    def test_empty_image_roots_cannot_silently_disable_http_path_policy(self):
        with (
            patch.dict("os.environ", {"GENPPT_IMAGE_ROOTS": ""}),
            self.assertRaisesRegex(ValueError, "GENPPT_IMAGE_ROOTS"),
        ):
            Settings.from_env()


if __name__ == "__main__":
    unittest.main()
